from __future__ import annotations

"""
Report formatters for different output formats
"""

import html
import json
import csv
import re
import sys
import tempfile
import os
import warnings
from typing import IO, Any, TextIO
from typing_extensions import override
from datetime import datetime
from pathlib import Path
import logging
from auto_a11y.models.test_result import ImpactLevel
from auto_a11y.reporting.comprehensive_report import ComprehensiveReportGenerator
from auto_a11y.reporting.issue_catalog import IssueCatalog
from io import StringIO

# Pattern to detect unresolved template placeholders in descriptions
_UNRESOLVED_PLACEHOLDER_RE = re.compile(r'\{[a-zA-Z_]+\}|%\([a-zA-Z_]+\)s')

logger = logging.getLogger(__name__)


# Hints we surface when WeasyPrint fails to load native libraries. These are
# the most common causes of macOS report-generation failures in production.
WEASYPRINT_NATIVE_LIB_HINTS = (
    'pango', 'cairo', 'gobject', 'gdk', 'libffi', 'glib', 'fontconfig',
    'harfbuzz', 'cannot load library', 'dlopen', 'image not found',
)


def platform_install_hint() -> str:
    """Return a platform-appropriate install hint for WeasyPrint native deps."""
    if sys.platform == 'darwin':
        return (
            "On macOS, install the required native libraries via Homebrew: "
            "brew install pango cairo gdk-pixbuf libffi"
        )
    if sys.platform.startswith('linux'):
        return (
            "On Debian/Ubuntu, install: "
            "apt-get install libpango-1.0-0 libpangoft2-1.0-0 libharfbuzz0b"
        )
    return (
        "Install WeasyPrint's native dependencies (Pango, Cairo, GDK-PixBuf). "
        "See https://doc.courtbouillon.org/weasyprint/stable/first_steps.html"
    )


def is_native_lib_error(err: BaseException) -> bool:
    """Heuristic: does this exception look like a missing native library?"""
    msg = str(err).lower()
    return any(hint in msg for hint in WEASYPRINT_NATIVE_LIB_HINTS)


def _dict_or_attr(obj: object, key: str, default: object = '') -> object:
    """Return obj[key] if obj is dict-like, else getattr(obj, key, default)."""
    if hasattr(obj, 'get'):
        return getattr(obj, 'get')(key, default)
    return getattr(obj, key, default)


class BaseFormatter:
    """Base class for report formatters"""

    TRANSLATIONS = {
        'en': {
            # Section headings
            'summary': 'Summary',
            'violations': 'Violations',
            'warnings': 'Warnings',
            'information_notes': 'Information Notes',
            'discovery_items': 'Discovery Items',
            'ai_analysis_findings': 'AI Analysis Findings',
            'passes': 'Passes',
            'overview': 'Overview',
            'projects_breakdown': 'Projects Breakdown',
            'violation_types': 'Error Types',
            'pages': 'Pages',
            'websites': 'Websites',
            'projects': 'Projects',
            # Report titles
            'accessibility_report': 'Accessibility Report',
            'all_projects_accessibility_report': 'All Projects Accessibility Report',
            'accessibility_executive_summary': 'Accessibility Executive Summary',
            'website_report': 'Website Report',
            'project_report': 'Project Report',
            'accessibility_report_summary': 'Accessibility Report Summary',
            # Labels
            'page': 'Page',
            'website': 'Website',
            'project': 'Project',
            'generated': 'Generated',
            'description': 'Description',
            'category': 'Category',
            'location': 'Location',
            'impact': 'Impact',
            'confidence': 'Confidence',
            'type': 'Type',
            'severity': 'Severity',
            'code': 'Code',
            'rule': 'Rule',
            'count': 'Count',
            'url': 'URL',
            'id': 'ID',
            'element': 'Element',
            'label': 'Label',
            'signature': 'Signature',
            'fix': 'Fix',
            # Compound labels
            'rule_id': 'Rule ID',
            'wcag_criteria': 'WCAG Criteria',
            'elements_affected': 'Elements Affected',
            'suggested_fix': 'Suggested Fix',
            'page_url': 'Page URL',
            'page_title': 'Page Title',
            'page_state': 'Page State',
            'session_id': 'Session ID',
            'test_duration': 'Test Duration',
            'total_projects': 'Total Projects',
            'total_websites': 'Total Websites',
            'total_pages': 'Total Pages',
            'pages_tested': 'Pages Tested',
            'total_violations': 'Total Violations',
            'total_warnings': 'Total Warnings',
            'pages_affected': 'Pages Affected',
            'last_tested': 'Last Tested',
            'test_date': 'Test Date',
            'website_name': 'Website Name',
            'project_name': 'Project Name',
            'tested_pages': 'Tested Pages',
            'coverage_pct': 'Coverage %',
            'breakpoint_label': 'Breakpoint',
            'css_state': 'CSS State',
            'touchpoint': 'Touchpoint',
            'xpath': 'XPath',
            'html': 'HTML',
            # Excel/sheet-specific
            'website_summary': 'Website Summary',
            'project_summary': 'Project Summary',
            'overall_summary': 'Overall Summary',
            'executive_summary': 'Executive Summary',
            'page_states': 'Page States',
            'all_issues': 'All Issues',
            'info': 'Info',
            'discovery': 'Discovery',
            'ai_findings': 'AI Findings',
            'all_issues_deduplicated': 'All Issues (Deduplicated)',
            'common_components': 'Common Components',
            'test_user': 'Test User',
            'user_roles': 'User Roles',
            'manual_check_required': 'Manual Check Required',
            'what': 'What',
            'why_important': 'Why Important',
            'who_affected': 'Who Affected',
            'how_to_remediate': 'How to Remediate',
            'location_xpath': 'Location (XPath)',
            'breakpoint_px': 'Breakpoint (px)',
            'pseudoclass': 'Pseudoclass',
            'component_type': 'Component Type',
            'page_count': 'Page Count',
            'pages_found': 'Pages Found',
            'example_xpath': 'Example XPath',
            'state_sequence': 'State Sequence',
            'state_description': 'State Description',
            'errors': 'Errors',
            'common_component_s': 'Common Component(s)',
            'pages_with_issue': 'Pages with Issue',
            'breakpoints': 'Breakpoints',
            'pseudoclasses': 'Pseudoclasses',
            'test_users': 'Test Users',
            # Stat card labels
            'passed': 'Passed',
            'errors_violations': 'Errors',
            # Phrases
            'manual_inspection_note': 'These items require manual inspection to ensure accessibility.',
            'no_websites_in_project': 'No websites in this project',
            'generated_by': 'Generated by Auto A11y Python',
            'generated_by_tool': 'Generated by Auto A11y - Accessibility Testing Tool',
            'report_generated_on': 'Report generated on',
            'unknown': 'Unknown',
            'no_description': 'No description',
            'not_specified': 'Not specified',
            'ai_finding': 'AI Finding',
            'violation': 'Error',
            'warning': 'Warning',
            'guest': 'Guest',
            'no_login': 'no login',
            # Impact levels
            'impact_high': 'High',
            'impact_medium': 'Medium',
            'impact_low': 'Low',
            'impact_info': 'Info',
            'impact_unknown': 'Unknown',
            # Summary sheet labels
            'summary_statistics': 'Summary Statistics',
            'overall_statistics': 'Overall Statistics',
            'average_violations_per_page': 'Average Errors/Page',
            'page_state_information': 'Page State Information',
            'state_details': 'State Details',
            'scripts_executed': 'Scripts Executed',
            'elements_clicked': 'Elements Clicked',
            'note': 'Note',
            'multistate_testing': 'Multi-State Testing',
            'multistate_note': 'This report shows results for a single page state. If multiple states were tested during the same session, each state will have its own test result with a different state_sequence value.',
            'breakpoint_testing': 'Breakpoint Testing',
            'breakpoint_note': "Some issues may have been detected at specific breakpoints (viewport widths). Check the 'Breakpoint (px)' column in the 'All Issues' sheet for breakpoint-specific issues.",
            'context_information': 'Context Information',
            'context_note': "Check the 'All Issues' sheet for complete context including Page State, Breakpoint, and Pseudoclass information for each issue.",
            # Executive report sections (streaming HTML report)
            'compliance': 'Compliance',
            'ai_executive_summary': 'AI Executive Analysis',
            'grade': 'Grade',
            'key_metrics': 'Key Metrics',
            'issues_by_touchpoint': 'Issues by Touchpoint',
            'issues_by_impact': 'Issues by Impact Level',
            'top_issues': 'Most Frequent Issues',
            'occurrences': 'Occurrences',
            'recommendations': 'Recommendations',
            'detailed_results': 'Detailed Results by Page',
            'average_page_score': 'Average Page Score',
            'compliance_score_desc': 'Share of automated checks that passed across all tested pages.',
            'exec_summary_text': '{pages} pages were tested: {violations} errors and {warnings} warnings were found, and {passes} automated checks passed.',
            'rec_high_impact': 'Address the {count} high-impact issues first — they most severely block users with disabilities.',
            'rec_touchpoint': 'Concentrate remediation on {name} — {count} issues, the largest category found.',
            'rec_warnings': 'Review the {count} warnings; many indicate real barriers that automated testing cannot fully confirm.',
            'rec_none': 'No errors were found. Continue periodic testing to maintain compliance.',
            'rec_retest': 'Re-test after each round of fixes to confirm remediation and catch regressions.',
        },
        'fr': {
            # Section headings
            'summary': 'Résumé',
            'violations': 'Violations',
            'warnings': 'Avertissements',
            'information_notes': "Notes d'information",
            'discovery_items': 'Éléments de découverte',
            'ai_analysis_findings': "Résultats de l'analyse IA",
            'passes': 'Réussites',
            'overview': "Vue d'ensemble",
            'projects_breakdown': 'Détail des projets',
            'violation_types': "Types d'erreurs",
            'pages': 'Pages',
            'websites': 'Sites web',
            'projects': 'Projets',
            # Report titles
            'accessibility_report': "Rapport d'accessibilité",
            'all_projects_accessibility_report': "Rapport d'accessibilité de tous les projets",
            'accessibility_executive_summary': "Résumé exécutif de l'accessibilité",
            'website_report': 'Rapport du site web',
            'project_report': 'Rapport du projet',
            'accessibility_report_summary': "Résumé du rapport d'accessibilité",
            # Labels
            'page': 'Page',
            'website': 'Site web',
            'project': 'Projet',
            'generated': 'Généré',
            'description': 'Description',
            'category': 'Catégorie',
            'location': 'Emplacement',
            'impact': 'Impact',
            'confidence': 'Confiance',
            'type': 'Type',
            'severity': 'Sévérité',
            'code': 'Code',
            'rule': 'Règle',
            'count': 'Nombre',
            'url': 'URL',
            'id': 'ID',
            'element': 'Élément',
            'label': 'Libellé',
            'signature': 'Signature',
            'fix': 'Correction',
            # Compound labels
            'rule_id': 'ID de règle',
            'wcag_criteria': 'Critères WCAG',
            'elements_affected': 'Éléments affectés',
            'suggested_fix': 'Correction suggérée',
            'page_url': 'URL de la page',
            'page_title': 'Titre de la page',
            'page_state': 'État de la page',
            'session_id': 'ID de session',
            'test_duration': 'Durée du test',
            'total_projects': 'Total des projets',
            'total_websites': 'Total des sites web',
            'total_pages': 'Total des pages',
            'pages_tested': 'Pages testées',
            'total_violations': 'Total des violations',
            'total_warnings': 'Total des avertissements',
            'pages_affected': 'Pages affectées',
            'last_tested': 'Dernier test',
            'test_date': 'Date de test',
            'website_name': 'Nom du site web',
            'project_name': 'Nom du projet',
            'tested_pages': 'Pages testées',
            'coverage_pct': 'Couverture %',
            'breakpoint_label': 'Point de rupture',
            'css_state': 'État CSS',
            'touchpoint': 'Point de contact',
            'xpath': 'XPath',
            'html': 'HTML',
            # Excel/sheet-specific
            'website_summary': 'Résumé du site web',
            'project_summary': 'Résumé du projet',
            'overall_summary': 'Résumé général',
            'executive_summary': 'Résumé exécutif',
            'page_states': 'États de la page',
            'all_issues': 'Tous les problèmes',
            'info': 'Infos',
            'discovery': 'Découverte',
            'ai_findings': 'Résultats IA',
            'all_issues_deduplicated': 'Problèmes (dédupliqués)',
            'common_components': 'Composants communs',
            'test_user': 'Utilisateur de test',
            'user_roles': 'Rôles utilisateur',
            'manual_check_required': 'Vérification manuelle requise',
            'what': 'Quoi',
            'why_important': "Pourquoi c'est important",
            'who_affected': 'Qui est affecté',
            'how_to_remediate': 'Comment corriger',
            'location_xpath': 'Emplacement (XPath)',
            'breakpoint_px': 'Point de rupture (px)',
            'pseudoclass': 'Pseudoclasse',
            'component_type': 'Type de composant',
            'page_count': 'Nombre de pages',
            'pages_found': 'Pages trouvées',
            'example_xpath': 'XPath exemple',
            'state_sequence': "Séquence d'état",
            'state_description': "Description de l'état",
            'errors': 'Erreurs',
            'common_component_s': 'Composant(s) commun(s)',
            'pages_with_issue': 'Pages avec problème',
            'breakpoints': 'Points de rupture',
            'pseudoclasses': 'Pseudoclasses',
            'test_users': 'Utilisateurs de test',
            # Stat card labels
            'passed': 'Réussites',
            'errors_violations': 'Erreurs',
            # Phrases
            'manual_inspection_note': "Ces éléments nécessitent une vérification manuelle pour assurer l'accessibilité.",
            'no_websites_in_project': 'Aucun site web dans ce projet',
            'generated_by': 'Généré par Auto A11y Python',
            'generated_by_tool': "Généré par Auto A11y - Outil de test d'accessibilité",
            'report_generated_on': 'Rapport généré le',
            'unknown': 'Inconnu',
            'no_description': 'Aucune description',
            'not_specified': 'Non spécifié',
            'ai_finding': 'Résultat IA',
            'violation': 'Erreur',
            'warning': 'Avertissement',
            'guest': 'Invité',
            'no_login': 'aucune connexion',
            # Impact levels
            'impact_high': 'Élevé',
            'impact_medium': 'Moyen',
            'impact_low': 'Faible',
            'impact_info': 'Info',
            'impact_unknown': 'Inconnu',
            # Summary sheet labels
            'summary_statistics': 'Statistiques résumées',
            'overall_statistics': 'Statistiques générales',
            'average_violations_per_page': 'Erreurs moyennes/page',
            'page_state_information': "Information sur l'état de la page",
            'state_details': "Détails de l'état",
            'scripts_executed': 'Scripts exécutés',
            'elements_clicked': 'Éléments cliqués',
            'note': 'Remarque',
            'multistate_testing': 'Test multi-états',
            'multistate_note': "Ce rapport affiche les résultats pour un seul état de page. Si plusieurs états ont été testés au cours de la même session, chaque état aura son propre résultat de test avec une valeur state_sequence différente.",
            'breakpoint_testing': 'Test de points de rupture',
            'breakpoint_note': "Certains problèmes peuvent avoir été détectés à des points de rupture spécifiques (largeurs de fenêtre). Vérifiez la colonne « Point de rupture (px) » dans la feuille « Tous les problèmes » pour les problèmes spécifiques aux points de rupture.",
            'context_information': 'Informations contextuelles',
            'context_note': "Consultez la feuille « Tous les problèmes » pour le contexte complet, y compris l'état de la page, le point de rupture et les informations de pseudoclasse pour chaque problème.",
            # Executive report sections (streaming HTML report)
            'compliance': 'Conformité',
            'ai_executive_summary': 'Analyse exécutive par IA',
            'grade': 'Note',
            'key_metrics': 'Indicateurs clés',
            'issues_by_touchpoint': 'Problèmes par point de contact',
            'issues_by_impact': "Problèmes par niveau d'impact",
            'top_issues': 'Problèmes les plus fréquents',
            'occurrences': 'Occurrences',
            'recommendations': 'Recommandations',
            'detailed_results': 'Résultats détaillés par page',
            'average_page_score': 'Score moyen des pages',
            'compliance_score_desc': 'Part des vérifications automatisées réussies sur l\'ensemble des pages testées.',
            'exec_summary_text': '{pages} pages ont été testées : {violations} erreurs et {warnings} avertissements ont été détectés, et {passes} vérifications automatisées ont réussi.',
            'rec_high_impact': "Corrigez d'abord les {count} problèmes à impact élevé — ce sont eux qui bloquent le plus gravement les personnes en situation de handicap.",
            'rec_touchpoint': 'Concentrez la remédiation sur {name} — {count} problèmes, la catégorie la plus importante détectée.',
            'rec_warnings': 'Examinez les {count} avertissements ; beaucoup signalent de véritables obstacles que les tests automatisés ne peuvent pas confirmer entièrement.',
            'rec_none': "Aucune erreur n'a été détectée. Poursuivez les tests périodiques pour maintenir la conformité.",
            'rec_retest': 'Testez à nouveau après chaque série de corrections pour confirmer la remédiation et détecter les régressions.',
            # Touchpoint display names
            'tp_accessible_names': 'Noms accessibles',
            'tp_animation': 'Animation',
            'tp_colors_contrast': 'Couleurs et contraste',
            'tp_color_contrast': 'Contraste des couleurs',
            'tp_colors': 'Couleurs',
            'tp_dialogs': 'Dialogues',
            'tp_electronic_documents': 'Documents électroniques',
            'tp_event_handling': 'Gestion des événements',
            'tp_fonts': 'Polices',
            'tp_forms': 'Formulaires',
            'tp_headings': 'Titres',
            'tp_images': 'Images',
            'tp_landmarks': 'Points de repère',
            'tp_language': 'Langue',
            'tp_links': 'Liens',
            'tp_lists': 'Listes',
            'tp_maps': 'Cartes',
            'tp_media': 'Multimédia',
            'tp_multimedia': 'Multimédia',
            'tp_navigation': 'Navigation',
            'tp_page': 'Page',
            'tp_semantic_structure': 'Structure sémantique',
            'tp_styles': 'Styles',
            'tp_tabindex': 'Index de tabulation',
            'tp_tables': 'Tableaux',
            'tp_timers': 'Minuteries',
            'tp_titles': 'Titres de page',
            'tp_focus_management': 'Gestion du focus',
            'tp_keyboard': 'Clavier',
            'tp_carousels': 'Carrousels',
            'tp_read_more_links': 'Liens « En savoir plus »',
            'tp_touch_and_gestures': 'Toucher et gestes',
            'tp_videos': 'Vidéos',
            'tp_buttons': 'Boutons',
            'tp_floating_content': 'Contenu flottant',
            'tp_title_attributes': 'Attributs de titre',
            'tp_pdf_annotations': 'Annotations PDF',
            'tp_pdf_document_properties': 'Propriétés du document PDF',
            'tp_pdf_tagging': 'Balisage PDF',
        },
    }

    def __init__(self, config: dict[str, Any], language: str = 'en') -> None:
        """Initialize formatter with config"""
        self.config = config
        self.language = language if language in ['en', 'fr'] else 'en'
        self.extension = 'txt'

    def _t(self, key: str) -> str:
        """Get translated string for current language"""
        return self.TRANSLATIONS.get(self.language, self.TRANSLATIONS['en']).get(key, key)

    def _translate_impact(self, impact_raw: str) -> str:
        """Translate an impact level value (high/medium/low) to the report language."""
        if not impact_raw:
            return ''
        key = f'impact_{impact_raw.strip().lower()}'
        translated = self._t(key)
        if translated != key:
            return translated.upper()
        # No specific translation for this impact value — fall back to the
        # translated 'unknown' label rather than leaking the raw English token
        # (which would otherwise appear untranslated in FR reports).
        return self._t('impact_unknown').upper()

    @staticmethod
    def _best_description(issue_dict: dict[str, Any]) -> str:
        """Pick the best available description from an enriched issue dict.

        Prefers description_full (instance-specific, with placeholders resolved),
        but falls back to what_generic (placeholder-free) when unresolved
        placeholders remain, and finally to the raw description field.
        """
        def _clean(text: str | None) -> bool:
            return bool(text and not _UNRESOLVED_PLACEHOLDER_RE.search(text))

        desc: str = issue_dict.get('description_full', '') or ''
        if _clean(desc):
            return desc
        generic: str = issue_dict.get('what_generic', '') or ''
        if _clean(generic):
            return generic
        raw: str = issue_dict.get('description', '') or ''
        if _clean(raw):
            return raw
        # The metadata dict may hold a clean English what_generic (stored by
        # result_processor at test time) and/or the original JS description
        # with values already interpolated.
        meta: dict[str, Any] = issue_dict.get('metadata') or {}
        meta_generic: str = str(meta.get('what_generic', '') or '')
        if _clean(meta_generic):
            return meta_generic
        meta_desc: str = str(meta.get('description', '') or '')
        if _clean(meta_desc):
            return meta_desc
        # Nothing clean available — return the best we have as-is
        return desc or generic or raw or ''

    @staticmethod
    def get_page_url(page: dict[str, Any] | object) -> str:
        """Extract page URL from page object or dict."""
        return str(_dict_or_attr(page, 'url', ''))

    @staticmethod
    def get_page_title(page: dict[str, Any] | object) -> str:
        """Extract page title from page object or dict."""
        return str(_dict_or_attr(page, 'title', ''))

    @staticmethod
    def get_issue_list(test_result: dict[str, Any] | object, attr: str) -> list[Any]:
        """Extract an issue list (violations/warnings) from test_result object or dict."""
        raw = _dict_or_attr(test_result, attr, [])
        if hasattr(raw, '__iter__') and not isinstance(raw, (str, bytes)):
            result: list[Any] = list(getattr(raw, '__iter__')())
            return result
        return []

    @staticmethod
    def to_enriched_dict(issue: dict[str, Any] | object) -> dict[str, Any]:
        """Convert an issue (object or dict) to an enriched dict."""
        raw: dict[str, Any]
        if hasattr(issue, 'items'):
            # dict-like object - use the attribute to get it as dict
            raw = dict(getattr(issue, 'items')())
        elif hasattr(issue, 'to_dict'):
            raw = getattr(issue, 'to_dict')()
        elif hasattr(issue, '__dict__'):
            raw = dict(vars(issue))
        else:
            raw = {}
        return IssueCatalog.enrich_issue(raw)

    @staticmethod
    def _format_wcag(issue_dict: dict[str, Any]) -> str:
        """Format WCAG criteria from an issue dict to a comma-separated string."""
        wcag_val: list[str] = issue_dict.get('wcag_criteria', [])
        return ', '.join(str(c) for c in wcag_val) if wcag_val else ''

    def format_page_report(self, data: dict[str, Any]) -> str | bytes:
        """Format page report data"""
        raise NotImplementedError
    
    def format_website_report(self, data: dict[str, Any]) -> str | bytes:
        """Format website report data"""
        warnings.warn(
            "format_website_report() is deprecated, use begin/append_page/finalize streaming interface",
            DeprecationWarning,
            stacklevel=2
        )
        raise NotImplementedError

    def format_project_report(self, data: dict[str, Any]) -> str | bytes:
        """Format project report data"""
        warnings.warn(
            "format_project_report() is deprecated, use begin/append_page/finalize streaming interface",
            DeprecationWarning,
            stacklevel=2
        )
        raise NotImplementedError
    
    def format_summary_report(self, data: dict[str, Any]) -> str | bytes:
        """Format summary report data"""
        raise NotImplementedError

    # --- Streaming interface ---

    def begin(self, output_file: str, summary: dict[str, Any]) -> None:
        """Write report header/preamble using summary stats."""
        raise NotImplementedError

    def append_page(self, output_file: str, page_data: dict[str, Any]) -> None:
        """Write one page's detail section."""
        raise NotImplementedError

    def finalize(self, output_file: str, summary: dict[str, Any]) -> None:
        """Write report footer/closing."""
        raise NotImplementedError

    def cleanup(self) -> None:
        """Delete any temp files created during streaming. Default no-op."""
        pass


class HTMLFormatter(BaseFormatter):
    """HTML report formatter"""

    # Upper bound on violation samples retained for AI-summary context;
    # keeps streaming memory O(1) regardless of site size.
    _AI_SAMPLE_CAP = 40

    def __init__(self, config: dict[str, Any], language: str = 'en') -> None:
        super().__init__(config, language)
        self.extension = 'html'
        # Pass Claude API key if available
        claude_api_key = config.get('CLAUDE_API_KEY')
        self.comprehensive_generator = ComprehensiveReportGenerator(claude_api_key=claude_api_key)
        # Streaming report shape. The standalone HTML report is an
        # executive summary (the offline/deduplicated reports carry the
        # per-page detail); the PDF formatter flips these when delegating
        # so its output keeps the detail section and skips the AI call.
        self.include_page_detail: bool = False
        self.include_ai_summary: bool = True
        # Streaming state (initialized in begin())
        self._output_file: str = ''
        self._summary: dict[str, Any] = {}
        self._body_tempfile: IO[str] | None = None
        self._ai_samples: list[dict[str, Any]] = []

    @staticmethod
    def _esc(value: object) -> str:
        """Escape a dynamic value for safe interpolation into HTML text.

        Report data can include page titles, URLs, descriptions, xpaths and
        HTML fragments captured from attacker-controlled tested sites. Every
        such value must be escaped to prevent stored XSS in generated reports.
        """
        return html.escape('' if value is None else str(value))

    @staticmethod
    def _esc_attr(value: object) -> str:
        """Escape a dynamic value for use inside a double-quoted HTML attribute."""
        return html.escape('' if value is None else str(value), quote=True)

    def format_all_projects_report(self, data: dict[str, Any]) -> str:
        """Format report for all projects as HTML"""
        warnings.warn(
            "format_all_projects_report() is deprecated, use begin/append_page/finalize streaming interface",
            DeprecationWarning,
            stacklevel=2
        )

        report_title = self._esc(data.get('title', self._t('all_projects_accessibility_report')))
        html = f"""<!DOCTYPE html>
<html lang="{self.language}">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{report_title}</title>
    {self._get_css()}
</head>
<body>
    <div class="container">
        <header>
            <h1>{report_title}</h1>
            <div class="metadata">
                <p><strong>{self._t('total_projects')}:</strong> {data['summary']['total_projects']}</p>
                <p><strong>{self._t('generated')}:</strong> {self._esc(data['generated_at'])}</p>
            </div>
        </header>

        <section class="summary">
            <h2>{self._t('overall_summary')}</h2>
            <div class="stats-grid">
                <div class="stat-card">
                    <h3>{data['summary']['total_projects']}</h3>
                    <p>{self._t('projects')}</p>
                </div>
                <div class="stat-card">
                    <h3>{data['summary']['total_websites']}</h3>
                    <p>{self._t('websites')}</p>
                </div>
                <div class="stat-card">
                    <h3>{data['summary']['total_pages']}</h3>
                    <p>{self._t('total_pages')}</p>
                </div>
                <div class="stat-card">
                    <h3>{data['summary']['total_tested']}</h3>
                    <p>{self._t('pages_tested')}</p>
                </div>
                <div class="stat-card violation">
                    <h3>{data['summary']['total_violations']}</h3>
                    <p>{self._t('total_violations')}</p>
                </div>
                <div class="stat-card warning">
                    <h3>{data['summary']['total_warnings']}</h3>
                    <p>{self._t('total_warnings')}</p>
                </div>
            </div>
        </section>

        <section class="projects">
            <h2>{self._t('projects_breakdown')}</h2>
            {"".join(self._format_project_section(p) for p in data['projects'])}
        </section>

        {self._get_footer()}
    </div>
</body>
</html>"""
        
        return html
    
    def _format_project_section(self, project_data: dict[str, Any]) -> str:
        """Format a single project section for all projects report"""
        project = project_data['project']
        stats = project_data['stats']
        
        websites_html = ""
        for website_data in project_data['websites']:
            website = website_data['website']
            websites_html += f"""
            <div class="website-item">
                <h4>{self._esc(website.get('name', website.get('url', self._t('unknown'))))}</h4>
                <p>{self._t('pages')}: {website_data['pages']} | {self._t('tested_pages')}: {website_data['tested']} |
                   {self._t('violations')}: {website_data['violations']} | {self._t('warnings')}: {website_data['warnings']}</p>
            </div>
            """

        return f"""
        <div class="project-section">
            <h3>{self._esc(project.get('name', self._t('unknown')))}</h3>
            <p>{self._esc(project.get('description', ''))}</p>
            <div class="project-stats">
                <span>{self._t('websites')}: {stats['website_count']}</span>
                <span>{self._t('pages')}: {stats['total_pages']}</span>
                <span>{self._t('tested_pages')}: {stats['tested_pages']}</span>
                <span>{self._t('coverage_pct')}: {stats['test_coverage']:.1f}%</span>
            </div>
            <div class="websites-list">
                {websites_html if websites_html else f'<p>{self._t("no_websites_in_project")}</p>'}
            </div>
        </div>
        """
    
    @override
    def format_page_report(self, data: dict[str, Any]) -> str:
        """Generate HTML report for a page"""

        # Check for page state information
        test_result = data.get('test_result', {})
        page_state_info = ""
        if test_result and test_result.get('session_id'):
            page_state: dict[str, Any] = test_result.get('page_state', {}) or {}
            state_desc: str = str(page_state.get('description', ''))

            if state_desc:
                page_state_info = f"""
                <p><strong>{self._t('page_state')}:</strong> {self._esc(state_desc)}</p>
                <p><strong>{self._t('session_id')}:</strong> {self._esc(test_result.get('session_id', ''))}</p>
                """

        page_url = data['page']['url']
        html = f"""<!DOCTYPE html>
<html lang="{self.language}">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{self._t('accessibility_report')} - {self._esc(page_url)}</title>
    {self._get_css()}
</head>
<body>
    <div class="container">
        <header>
            <h1>{self._t('accessibility_report')}</h1>
            <div class="metadata">
                <p><strong>{self._t('page')}:</strong> <a href="{self._esc_attr(page_url)}" target="_blank">{self._esc(page_url)}</a></p>
                <p><strong>{self._t('website')}:</strong> {self._esc(data['website']['name'])}</p>
                <p><strong>{self._t('project')}:</strong> {self._esc(data['project']['name'])}</p>
                <p><strong>{self._t('generated')}:</strong> {self._esc(data['generated_at'])}</p>
                {page_state_info}
            </div>
        </header>

        {self._format_summary_section(data['statistics'])}

        {self._format_violations_section(data.get('violations', []))}

        {self._format_warnings_section(data.get('warnings', []))}

        {self._format_info_section(data.get('info', []))}

        {self._format_discovery_section(data.get('discovery', []))}

        {self._format_ai_findings_section(data.get('ai_findings', []))}

        {self._format_passes_section(data.get('passes', []))}

        <footer>
            <p>{self._t('generated_by')} - {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</p>
        </footer>
    </div>
</body>
</html>"""

        return html
    
    @override
    def format_website_report(self, data: dict[str, Any]) -> str:
        """Generate HTML report for a website"""

        # Use comprehensive report generator for website reports
        return self.comprehensive_generator.generate_comprehensive_html(data)

    @override
    def format_project_report(self, data: dict[str, Any]) -> str:
        """Generate bilingual HTML report for a project"""

        # Generate temporary output path
        project_name = data.get('project', {}).get('name', 'project').replace(' ', '_')
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        output_path = os.path.join(tempfile.gettempdir(), f"project_{project_name}_{timestamp}.html")

        # Use new bilingual standalone report generator
        self.comprehensive_generator.generate_bilingual_standalone_html(data, output_path, include_ai_summary=True)

        # Read and return the generated HTML
        with open(output_path, 'r', encoding='utf-8') as f:
            return f.read()
    
    @override
    def format_summary_report(self, data: dict[str, Any]) -> str:
        """Generate executive summary report"""

        projects_html = ""
        for project in data['projects']:
            projects_html += f"""
            <tr>
                <td>{self._esc(project['name'])}</td>
                <td>{project['websites']}</td>
                <td>{project['pages_tested']}</td>
                <td class="violations">{project['violations']}</td>
            </tr>"""
        
        html = f"""<!DOCTYPE html>
<html lang="{self.language}">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{self._t('accessibility_executive_summary')}</title>
    {self._get_css()}
</head>
<body>
    <div class="container">
        <header>
            <h1>{self._t('accessibility_executive_summary')}</h1>
            <p class="generated">{self._t('generated')}: {self._esc(data['generated_at'])}</p>
        </header>

        <section class="overview">
            <h2>{self._t('overview')}</h2>
            <div class="stats-grid">
                <div class="stat-card">
                    <h3>{len(data['projects'])}</h3>
                    <p>{self._t('projects')}</p>
                </div>
                <div class="stat-card">
                    <h3>{data['total_pages_tested']}</h3>
                    <p>{self._t('pages_tested')}</p>
                </div>
                <div class="stat-card violations">
                    <h3>{data['total_violations']}</h3>
                    <p>{self._t('total_violations')}</p>
                </div>
            </div>
        </section>

        <section class="projects">
            <h2>{self._t('projects')}</h2>
            <table>
                <thead>
                    <tr>
                        <th>{self._t('project')}</th>
                        <th>{self._t('websites')}</th>
                        <th>{self._t('pages_tested')}</th>
                        <th>{self._t('violations')}</th>
                    </tr>
                </thead>
                <tbody>
                    {projects_html}
                </tbody>
            </table>
        </section>

        <footer>
            <p>{self._t('generated_by')} - {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</p>
        </footer>
    </div>
</body>
</html>"""
        
        return html
    
    def _get_footer(self) -> str:
        """Get footer for HTML reports"""
        return f"""
        <footer>
            <p>{self._t('generated_by_tool')}</p>
            <p>{self._t('report_generated_on')} {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</p>
        </footer>
        """
    
    def _get_css(self) -> str:
        """Get CSS styles for HTML reports"""
        # Standalone reports cannot link the app stylesheets, so severity/brand
        # colours mirror the light-mode design tokens in
        # auto_a11y/web/static/public/css/tokens.css. Every text colour below
        # meets WCAG 2.2 AA (4.5:1+ on its background; 3:1+ for large text and
        # non-text indicators) — keep it that way when editing.
        return """
    <style>
        * { margin: 0; padding: 0; box-sizing: border-box; }
        body {
            font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Arial, sans-serif;
            line-height: 1.6;
            color: #333;
            background: #f5f5f5;
        }
        .container {
            max-width: 1200px;
            margin: 0 auto;
            padding: 20px;
            background: white;
            box-shadow: 0 0 10px rgba(0,0,0,0.1);
        }
        header {
            border-bottom: 2px solid #1a5276;
            padding-bottom: 20px;
            margin-bottom: 30px;
        }
        h1 { color: #1a5276; margin-bottom: 10px; }
        h2 { color: #333; margin: 20px 0 15px; border-bottom: 1px solid #e0e0e0; padding-bottom: 5px; }
        h3 { color: #555; margin: 15px 0 10px; }
        
        .metadata {
            background: #f8f9fa;
            padding: 15px;
            border-radius: 5px;
            margin: 15px 0;
        }
        .metadata p { margin: 5px 0; }
        
        .stats-grid {
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(200px, 1fr));
            gap: 20px;
            margin: 20px 0;
        }
        
        .stat-card {
            background: #f8f9fa;
            padding: 20px;
            border-radius: 8px;
            text-align: center;
            border: 1px solid #dee2e6;
        }
        .stat-card h3 {
            font-size: 2em;
            color: #1a5276;
            margin: 0;
        }
        .stat-card.violations h3 { color: #922b21; }
        .stat-card.warnings h3 { color: #7d6608; }
        .stat-card.passes h3 { color: #1a7a3e; }

        .violation, .warning, .pass, .ai-finding {
            background: #fff;
            border-left: 4px solid #922b21;
            padding: 15px;
            margin: 15px 0;
            border-radius: 4px;
            box-shadow: 0 2px 4px rgba(0,0,0,0.1);
        }
        .warning { border-left-color: #b38600; }
        .pass { border-left-color: #1a7a3e; }
        .ai-finding { border-left-color: #6f42c1; }
        
        .violation h4, .warning h4, .pass h4, .ai-finding h4 {
            margin-bottom: 10px;
            color: #333;
        }
        
        .impact {
            display: inline-block;
            padding: 2px 8px;
            border-radius: 3px;
            font-size: 0.85em;
            font-weight: bold;
            margin-left: 10px;
        }
        .impact.critical { background: #922b21; color: white; }
        .impact.serious { background: #b85700; color: white; }
        .impact.moderate { background: #ffc107; color: #333; }
        .impact.minor { background: #6c757d; color: white; }
        
        table {
            width: 100%;
            border-collapse: collapse;
            margin: 20px 0;
        }
        th, td {
            padding: 12px;
            text-align: left;
            border-bottom: 1px solid #dee2e6;
        }
        th {
            background: #f8f9fa;
            font-weight: bold;
            color: #495057;
        }
        tr:hover { background: #f8f9fa; }
        
        .code {
            background: #f4f4f4;
            padding: 10px;
            border-radius: 4px;
            font-family: 'Courier New', monospace;
            font-size: 0.9em;
            overflow-x: auto;
            margin: 10px 0;
        }
        
        footer {
            margin-top: 40px;
            padding-top: 20px;
            border-top: 1px solid #dee2e6;
            text-align: center;
            color: #6c757d;
        }
        
        a { color: #0a58ca; text-decoration: none; }
        a:hover { text-decoration: underline; }
    </style>
        """
    
    def _format_summary_section(self, stats: dict[str, Any]) -> str:
        """Format summary statistics section"""
        return f"""
        <section class="summary">
            <h2>{self._t('summary')}</h2>
            <div class="stats-grid">
                <div class="stat-card violations">
                    <h3>{stats['violations']}</h3>
                    <p>{self._t('violations')}</p>
                </div>
                <div class="stat-card warnings">
                    <h3>{stats['warnings']}</h3>
                    <p>{self._t('warnings')}</p>
                </div>
                <div class="stat-card passes">
                    <h3>{stats['passes']}</h3>
                    <p>{self._t('passes')}</p>
                </div>
                <div class="stat-card">
                    <h3>{stats['duration_ms']}ms</h3>
                    <p>{self._t('test_duration')}</p>
                </div>
            </div>
        </section>"""
    
    def _format_violations_section(self, violations: list[dict[str, Any]]) -> str:
        """Format violations section"""
        if not violations:
            return ""

        html = f"<section class='violations'><h2>{self._t('violations')}</h2>"
        for v in violations:
            impact_class = v.get('impact', 'moderate').lower()
            metadata = v.get('metadata', {})

            # Build metadata display
            metadata_html = ""
            if metadata.get('breakpoint'):
                metadata_html += f"<p><strong>{self._t('breakpoint_label')}:</strong> {self._esc(metadata.get('breakpoint'))}px</p>"
            if metadata.get('pseudoclass'):
                metadata_html += f"<p><strong>{self._t('css_state')}:</strong> {self._esc(metadata.get('pseudoclass'))}</p>"

            html += f"""
            <div class="violation">
                <h4>{self._esc(v.get('rule_id', self._t('unknown')))}
                    <span class="impact {self._esc_attr(impact_class)}">{self._translate_impact(v.get('impact', 'moderate'))}</span>
                </h4>
                <p><strong>{self._t('description')}:</strong> {self._esc(self._best_description(v))}</p>
                <p><strong>{self._t('wcag_criteria')}:</strong> {self._esc(', '.join(v.get('wcag_criteria', [])))}</p>
                <p><strong>{self._t('elements_affected')}:</strong> {v.get('node_count', 0)}</p>
                {metadata_html}
                {self._format_fix(v.get('suggested_fix'))}
            </div>"""
        html += "</section>"
        return html
    
    def _format_warnings_section(self, warnings: list[dict[str, Any]]) -> str:
        """Format warnings section"""
        if not warnings:
            return ""

        html = f"<section class='warnings'><h2>{self._t('warnings')}</h2>"
        for w in warnings:
            metadata = w.get('metadata', {})

            # Build metadata display
            metadata_html = ""
            if metadata.get('breakpoint'):
                metadata_html += f"<p><strong>{self._t('breakpoint_label')}:</strong> {self._esc(metadata.get('breakpoint'))}px</p>"
            if metadata.get('pseudoclass'):
                metadata_html += f"<p><strong>{self._t('css_state')}:</strong> {self._esc(metadata.get('pseudoclass'))}</p>"

            html += f"""
            <div class="warning">
                <h4>{self._esc(w.get('rule_id', self._t('unknown')))}</h4>
                <p>{self._esc(self._best_description(w))}</p>
                {metadata_html}
            </div>"""
        html += "</section>"
        return html
    
    def _format_info_section(self, info_items: list[dict[str, Any]]) -> str:
        """Format info section"""
        if not info_items:
            return ""

        html = f"<section class='info'><h2>{self._t('information_notes')}</h2>"
        for item in info_items:
            metadata = item.get('metadata', {})

            # Build metadata display
            metadata_html = ""
            if metadata.get('breakpoint'):
                metadata_html += f"<p><strong>{self._t('breakpoint_label')}:</strong> {self._esc(metadata.get('breakpoint'))}px</p>"
            if metadata.get('pseudoclass'):
                metadata_html += f"<p><strong>{self._t('css_state')}:</strong> {self._esc(metadata.get('pseudoclass'))}</p>"

            html += f"""
            <div class="info-item">
                <h4>{self._esc(item.get('id', self._t('unknown')))}</h4>
                <p><strong>{self._t('description')}:</strong> {self._esc(self._best_description(item))}</p>
                <p><strong>{self._t('category')}:</strong> {self._esc(item.get('category', 'General'))}</p>
                {f"<p><strong>{self._t('wcag_criteria')}:</strong> {self._esc(', '.join(item.get('wcag_criteria', [])))}</p>" if item.get('wcag_criteria') else ""}
                {metadata_html}
            </div>"""
        html += "</section>"
        return html

    def _format_discovery_section(self, discovery_items: list[dict[str, Any]]) -> str:
        """Format discovery section"""
        if not discovery_items:
            return ""

        html = f"<section class='discovery'><h2>{self._t('discovery_items')}</h2>"
        html += f"<p><em>{self._t('manual_inspection_note')}</em></p>"
        for item in discovery_items:
            metadata = item.get('metadata', {})

            # Build metadata display
            metadata_html = ""
            if metadata.get('breakpoint'):
                metadata_html += f"<p><strong>{self._t('breakpoint_label')}:</strong> {self._esc(metadata.get('breakpoint'))}px</p>"
            if metadata.get('pseudoclass'):
                metadata_html += f"<p><strong>{self._t('css_state')}:</strong> {self._esc(metadata.get('pseudoclass'))}</p>"

            html += f"""
            <div class="discovery-item">
                <h4>{self._esc(item.get('id', self._t('unknown')))}</h4>
                <p><strong>{self._t('description')}:</strong> {self._esc(self._best_description(item))}</p>
                <p><strong>{self._t('category')}:</strong> {self._esc(item.get('category', 'General'))}</p>
                {f"<p><strong>{self._t('location')}:</strong> <code>{self._esc(item.get('xpath', self._t('not_specified')))}</code></p>" if item.get('xpath') else ""}
                {metadata_html}
            </div>"""
        html += "</section>"
        return html

    def _format_ai_findings_section(self, findings: list[Any]) -> str:
        """Format AI findings section"""
        if not findings:
            return ""
        
        html = f"<section class='ai-findings'><h2>{self._t('ai_analysis_findings')}</h2>"
        for f in findings:
            severity_class = f.severity.value.lower() if hasattr(f, 'severity') else 'moderate'
            html += f"""
            <div class="ai-finding">
                <h4>{self._esc(f.type if hasattr(f, 'type') else self._t('ai_finding'))}
                    <span class="impact {self._esc_attr(severity_class)}">{f.severity.value.upper() if hasattr(f, 'severity') else 'MODERATE'}</span>
                </h4>
                <p><strong>{self._t('description')}:</strong> {self._esc(f.description if hasattr(f, 'description') else self._t('no_description'))}</p>
                <p><strong>{self._t('confidence')}:</strong> {f.confidence * 100 if hasattr(f, 'confidence') else 85:.0f}%</p>
                {self._format_fix(f.suggested_fix if hasattr(f, 'suggested_fix') else None)}
            </div>"""
        html += "</section>"
        return html
    
    def _format_passes_section(self, passes: list[dict[str, Any]]) -> str:
        """Format passes section"""
        if not passes:
            return ""
        
        html = f"<section class='passes'><h2>{self._t('passes')}</h2><ul>"
        for p in passes[:10]:  # Show first 10 passes
            html += f"<li>{self._esc(p.get('rule_id', self._t('unknown')))}: {self._esc(p.get('description', self._t('passed')))}</li>"
        if len(passes) > 10:
            remaining = len(passes) - 10
            html += f"<li>... {remaining} more</li>"
        html += "</ul></section>"
        return html
    
    def _format_fix(self, fix: str | None) -> str:
        """Format suggested fix"""
        if not fix:
            return ""
        return f'<p><strong>{self._t("suggested_fix")}:</strong> {self._esc(fix)}</p>'
    
    def _format_violation_types_section(self, violation_types: dict[str, Any]) -> str:
        """Format violation types section"""
        html = f"<section class='violation-types'><h2>{self._t('violation_types')}</h2><table>"
        html += f"<thead><tr><th>{self._t('rule')}</th><th>{self._t('count')}</th><th>{self._t('description')}</th><th>{self._t('pages_affected')}</th></tr></thead><tbody>"
        
        sorted_types = sorted(violation_types.items(), key=lambda x: x[1]['count'], reverse=True)
        for rule_id, info in sorted_types[:20]:  # Show top 20
            pages_summary = f"{len(set(info['pages']))} pages"
            html += f"""
            <tr>
                <td>{self._esc(rule_id)}</td>
                <td>{info['count']}</td>
                <td>{self._esc(info['description'])}</td>
                <td>{self._esc(pages_summary)}</td>
            </tr>"""
        
        html += "</tbody></table></section>"
        return html
    
    def _format_pages_table(self, page_results: list[dict[str, Any]]) -> str:
        """Format pages table"""
        html = f"<section class='pages'><h2>{self._t('pages')}</h2><table>"
        html += f"<thead><tr><th>{self._t('page')}</th><th>{self._t('violations')}</th><th>{self._t('warnings')}</th><th>{self._t('last_tested')}</th></tr></thead><tbody>"
        
        for pr in page_results:
            page = pr['page']
            test = pr['test_result']
            html += f"""
            <tr>
                <td><a href="{self._esc_attr(page.url)}" target="_blank">{self._esc(page.url)}</a></td>
                <td class="violations">{test.violation_count}</td>
                <td class="warnings">{test.warning_count}</td>
                <td>{self._esc(test.test_date)}</td>
            </tr>"""
        
        html += "</tbody></table></section>"
        return html
    
    def _format_websites_section(self, websites: list[dict[str, Any]]) -> str:
        """Format websites section with detailed test results"""
        html = f"<section class='websites'><h2>{self._t('websites')}</h2>"
        
        for wd in websites:
            website = wd['website']
            pages = wd['pages']
            
            # Calculate totals for all categories
            total_violations = 0
            total_warnings = 0
            total_info = 0
            total_discovery = 0
            total_passes = 0
            
            all_violations: list[dict[str, Any]] = []
            all_warnings: list[dict[str, Any]] = []
            all_info: list[dict[str, Any]] = []
            all_discovery: list[dict[str, Any]] = []

            for p in pages:
                test_result = p['test_result']
                if hasattr(test_result, 'violation_count'):
                    total_violations += test_result.violation_count
                    total_warnings += test_result.warning_count
                    total_info += test_result.info_count
                    total_discovery += test_result.discovery_count
                    total_passes += test_result.pass_count

                    # Collect all issues for detailed display
                    for v in test_result.violations:
                        all_violations.append({'page': p['page'].url, 'issue': v})
                    for w in test_result.warnings:
                        all_warnings.append({'page': p['page'].url, 'issue': w})
                    for i_item in test_result.info:
                        all_info.append({'page': p['page'].url, 'issue': i_item})
                    for d in test_result.discovery:
                        all_discovery.append({'page': p['page'].url, 'issue': d})
            
            # Website header with all statistics
            html += f"""
            <div class="website">
                <h3>{self._esc(website.name)}</h3>
                <p><a href="{self._esc_attr(website.url)}" target="_blank">{self._esc(website.url)}</a></p>
                <div class="stats-grid" style="margin: 20px 0;">
                    <div class="stat-card violations">
                        <h4>{total_violations}</h4>
                        <p>{self._t('errors')}</p>
                    </div>
                    <div class="stat-card warnings">
                        <h4>{total_warnings}</h4>
                        <p>{self._t('warnings')}</p>
                    </div>
                    <div class="stat-card info">
                        <h4>{total_info}</h4>
                        <p>{self._t('info')}</p>
                    </div>
                    <div class="stat-card discovery">
                        <h4>{total_discovery}</h4>
                        <p>{self._t('discovery')}</p>
                    </div>
                    <div class="stat-card passes">
                        <h4>{total_passes}</h4>
                        <p>{self._t('passed')}</p>
                    </div>
                </div>

                <p><strong>{self._t('pages_tested')}:</strong> {len(pages)}</p>
                """
            
            # Add detailed issues if they exist
            if all_violations:
                html += f"<h4>{self._t('errors_violations')}</h4><ul>"
                for v_item in all_violations[:10]:  # Show first 10
                    v_issue: object = v_item['issue']
                    issue_id: str = str(getattr(v_issue, 'id', self._t('unknown')))
                    issue_desc: str = str(getattr(v_issue, 'description', self._t('no_description')))
                    html += f"<li><strong>{self._esc(issue_id)}:</strong> {self._esc(issue_desc)} - <em>{self._esc(v_item['page'])}</em></li>"
                if len(all_violations) > 10:
                    remaining = len(all_violations) - 10
                    html += f"<li><em>... {remaining} more</em></li>"
                html += "</ul>"

            if all_warnings:
                html += f"<h4>{self._t('warnings')}</h4><ul>"
                for w_item in all_warnings[:10]:  # Show first 10
                    w_issue: object = w_item['issue']
                    issue_id = str(getattr(w_issue, 'id', self._t('unknown')))
                    issue_desc = str(getattr(w_issue, 'description', self._t('no_description')))
                    html += f"<li><strong>{self._esc(issue_id)}:</strong> {self._esc(issue_desc)} - <em>{self._esc(w_item['page'])}</em></li>"
                if len(all_warnings) > 10:
                    remaining = len(all_warnings) - 10
                    html += f"<li><em>... {remaining} more</em></li>"
                html += "</ul>"

            if all_info:
                html += f"<h4>{self._t('information_notes')}</h4><ul>"
                for info_item in all_info[:5]:  # Show first 5
                    i_issue: object = info_item['issue']
                    issue_id = str(getattr(i_issue, 'id', self._t('unknown')))
                    issue_desc = str(getattr(i_issue, 'description', self._t('no_description')))
                    html += f"<li><strong>{self._esc(issue_id)}:</strong> {self._esc(issue_desc)} - <em>{self._esc(info_item['page'])}</em></li>"
                if len(all_info) > 5:
                    remaining = len(all_info) - 5
                    html += f"<li><em>... {remaining} more</em></li>"
                html += "</ul>"

            if all_discovery:
                html += f"<h4>{self._t('discovery_items')}</h4><ul>"
                for d_item in all_discovery[:5]:  # Show first 5
                    d_issue: object = d_item['issue']
                    issue_id = str(getattr(d_issue, 'id', self._t('unknown')))
                    issue_desc = str(getattr(d_issue, 'description', self._t('no_description')))
                    html += f"<li><strong>{self._esc(issue_id)}:</strong> {self._esc(issue_desc)} - <em>{self._esc(d_item['page'])}</em></li>"
                if len(all_discovery) > 5:
                    remaining = len(all_discovery) - 5
                    html += f"<li><em>... {remaining} more</em></li>"
                html += "</ul>"
            
            html += "</div>"
        
        html += "</section>"
        return html

    # --- Streaming interface ---

    @override
    def begin(self, output_file: str, summary: dict[str, Any]) -> None:
        """Prepare streaming state.

        The final output is NOT written yet. When per-page detail is
        included (PDF delegation), a temp body file is created so that
        ``append_page`` can write into it; the standalone HTML report is
        executive-only, so no body file is needed there.
        """
        self._output_file = output_file
        self._summary = dict(summary) if summary else {}
        self._ai_samples = []
        if self.include_page_detail:
            self._body_tempfile = tempfile.NamedTemporaryFile(
                mode='w', suffix='.html', delete=False, encoding='utf-8'
            )
        else:
            self._body_tempfile = None

    def _collect_ai_samples(self, violations: list[Any]) -> None:
        """Accumulate a bounded sample of violations as AI-summary context.

        High-impact issues are preferred; the cap keeps memory O(1) no
        matter how many pages stream through.
        """
        for v in violations:
            if len(self._ai_samples) >= self._AI_SAMPLE_CAP:
                high = [x for x in self._ai_samples if x.get('impact') == 'high']
                if len(high) >= self._AI_SAMPLE_CAP:
                    return
            issue = self.to_enriched_dict(v)
            impact_raw = issue.get('impact', '')
            if hasattr(impact_raw, 'value'):
                impact_raw = impact_raw.value
            impact = str(impact_raw).lower()
            sample = {
                'description': self._best_description(issue),
                'impact': impact,
                'wcag': issue.get('wcag_criteria', []),
                'category': str(issue.get('touchpoint', 'general')),
            }
            if len(self._ai_samples) < self._AI_SAMPLE_CAP:
                self._ai_samples.append(sample)
            elif impact == 'high':
                # Displace the first non-high sample to keep the most
                # severe issues in the window.
                for idx, existing in enumerate(self._ai_samples):
                    if existing.get('impact') != 'high':
                        self._ai_samples[idx] = sample
                        break

    @override
    def append_page(self, output_file: str, page_data: dict[str, Any]) -> None:
        """Stream one page: collect AI context and (when enabled) write its
        violations/warnings as an HTML section to the temp body file."""
        page = page_data.get('page', {})
        test_result = page_data.get('test_result')
        if test_result is None:
            return

        page_url: str = self.get_page_url(page)
        page_title: str = self.get_page_title(page)

        violations: list[Any] = self.get_issue_list(test_result, 'violations')
        warnings_list: list[Any] = self.get_issue_list(test_result, 'warnings')

        if self.include_ai_summary:
            self._collect_ai_samples(violations)

        if not self.include_page_detail:
            return

        section = f'<div class="page-section"><h3><a href="{self._esc_attr(page_url)}">{self._esc(page_title or page_url)}</a></h3>\n'

        if violations:
            section += f'<table><thead><tr><th>{self._t("code")}</th><th>{self._t("description")}</th><th>{self._t("impact")}</th><th>{self._t("wcag_criteria")}</th><th>{self._t("xpath")}</th></tr></thead><tbody>\n'
            for v in violations:
                section += self._streaming_issue_row(v)
            section += '</tbody></table>\n'

        if warnings_list:
            section += f'<h4>{self._t("warnings")}</h4>\n'
            section += f'<table><thead><tr><th>{self._t("code")}</th><th>{self._t("description")}</th><th>{self._t("impact")}</th><th>{self._t("wcag_criteria")}</th><th>{self._t("xpath")}</th></tr></thead><tbody>\n'
            for w in warnings_list:
                section += self._streaming_issue_row(w)
            section += '</tbody></table>\n'

        section += '</div>\n'
        if self._body_tempfile is not None:
            self._body_tempfile.write(section)

    @override
    def finalize(self, output_file: str, summary: dict[str, Any]) -> None:
        """Assemble the final HTML document as a structured executive report.

        The Pass-1 summary provides everything the executive sections need
        (totals, touchpoint/impact distributions, page scores, top issue
        codes); the per-page detail streamed by ``append_page`` is read back
        from the temp file in 64 KB chunks so the full body is never loaded
        into memory at once.
        """
        # Flush and close the temp body file so we can read it back
        if hasattr(self, '_body_tempfile') and self._body_tempfile and not self._body_tempfile.closed:
            self._body_tempfile.flush()
            self._body_tempfile.close()

        s = summary if summary else {}

        with open(output_file, 'w', encoding='utf-8') as out:
            out.write(f'<!DOCTYPE html>\n<html lang="{self.language}">\n<head>\n')
            out.write('<meta charset="UTF-8">\n')
            out.write('<meta name="viewport" content="width=device-width, initial-scale=1.0">\n')
            out.write(f'<title>{self._report_title(s)}</title>\n')
            out.write(self._get_css())
            out.write(self._get_exec_css())
            out.write('\n</head>\n<body>\n<div class="container">\n')
            out.write(self._render_report_header(s))
            out.write(self._render_executive_summary(s))
            out.write(self._render_key_metrics(s))
            out.write(self._render_touchpoint_chart(s))
            out.write(self._render_impact_chart(s))
            out.write(self._render_top_issues(s))
            out.write(self._render_recommendations(s))

            if self.include_ai_summary:
                out.write(self._render_ai_summary(s))

            if self.include_page_detail:
                out.write(f'<section class="detailed-results"><h2>{self._t("detailed_results")}</h2>\n')
                # Stream body from temp file in 64KB chunks
                body_path = getattr(self, '_body_tempfile', None)
                if body_path and hasattr(body_path, 'name') and os.path.exists(body_path.name):
                    with open(body_path.name, 'r', encoding='utf-8') as body:
                        while True:
                            chunk = body.read(65536)
                            if not chunk:
                                break
                            out.write(chunk)
                out.write('</section>\n')

            out.write(self._get_footer())
            out.write('\n</div>\n</body>\n</html>')

    def _render_ai_summary(self, s: dict[str, Any]) -> str:
        """Render the AI executive analysis from Pass-1 stats plus the
        bounded violation sample collected while streaming.

        Any failure (no API key, network error) falls back to the
        generator's deterministic summary; a hard failure skips the
        section rather than aborting the report.
        """
        from auto_a11y.reporting.ai_executive_summary import AIExecutiveSummaryGenerator

        generator: AIExecutiveSummaryGenerator
        claude_api_key = self.config.get('CLAUDE_API_KEY')
        if claude_api_key:
            try:
                from auto_a11y.ai.claude_client import ClaudeClient, ClaudeConfig
                claude_config = ClaudeConfig(
                    api_key=str(claude_api_key),
                    model=str(self.config.get('CLAUDE_MODEL') or 'claude-opus-4-8'),
                )
                generator = AIExecutiveSummaryGenerator(ClaudeClient(config=claude_config))
            except Exception as e:
                logger.warning(f"Could not initialize AI summary client: {e}")
                generator = AIExecutiveSummaryGenerator()
        else:
            generator = AIExecutiveSummaryGenerator()

        report_data: dict[str, Any] = {
            'project': s.get('project', {}),
            'statistics': {
                'total_pages': int(s.get('total_pages', 0) or 0),
                'total_violations': int(s.get('total_violations', 0) or 0),
                'total_warnings': int(s.get('total_warnings', 0) or 0),
                'total_info': int(s.get('total_info', 0) or 0),
                'total_discovery': int(s.get('total_discovery', 0) or 0),
                'total_passes': int(s.get('total_passes', 0) or 0),
            },
            'violation_samples': list(self._ai_samples),
            'recordings': s.get('recordings', []) or [],
        }

        try:
            ai_summary = generator.generate_executive_summary(report_data, self.language)
            body = generator.format_executive_summary_html(ai_summary, self.language)
            css = generator.get_ai_summary_css()
            return (
                f'<section class="ai-summary"><h2>{self._t("ai_executive_summary")}</h2>\n'
                f'<style>{css}</style>\n{body}\n</section>\n'
            )
        except Exception as e:
            logger.error(f"AI executive summary rendering failed: {e}", exc_info=True)
            return ''

    @override
    def cleanup(self) -> None:
        """Delete the temp body file."""
        if hasattr(self, '_body_tempfile') and self._body_tempfile:
            path = self._body_tempfile.name
            if not self._body_tempfile.closed:
                self._body_tempfile.close()
            if os.path.exists(path):
                os.unlink(path)

    # --- streaming helpers ---

    def _streaming_issue_row(self, issue: Any) -> str:
        """Return an HTML <tr> for one issue."""
        issue_dict: dict[str, Any] = self.to_enriched_dict(issue)

        impact_raw = issue_dict.get('impact', '')
        if hasattr(impact_raw, 'value'):
            impact_raw = impact_raw.value
        impact_str = self._translate_impact(str(impact_raw))

        wcag_str: str = self._format_wcag(issue_dict)

        code: str = str(issue_dict.get('id', ''))
        description = self._best_description(issue_dict)
        xpath: str = str(issue_dict.get('xpath', ''))

        return f'<tr><td>{self._esc(code)}</td><td>{self._esc(description)}</td><td>{impact_str}</td><td>{self._esc(wcag_str)}</td><td>{self._esc(xpath)}</td></tr>\n'

    # --- executive report sections (HTML streaming report only) ---

    def _report_title(self, s: dict[str, Any]) -> str:
        """Title for the report: explicit title, else scope name + report label."""
        explicit = s.get('title')
        if explicit:
            return self._esc(explicit)
        scope = s.get('project') or s.get('website') or {}
        name = scope.get('name') if isinstance(scope, dict) else None
        if name:
            return f'{self._esc(name)} — {self._t("accessibility_report")}'
        return self._t('accessibility_report')

    def _render_report_header(self, s: dict[str, Any]) -> str:
        generated = datetime.now().strftime('%Y-%m-%d %H:%M')
        subtitle_parts: list[str] = []
        project = s.get('project')
        if isinstance(project, dict) and project.get('name'):
            subtitle_parts.append(f'{self._t("project")}: {self._esc(project["name"])}')
        website = s.get('website')
        if isinstance(website, dict) and website.get('name'):
            subtitle_parts.append(f'{self._t("website")}: {self._esc(website["name"])}')
        subtitle = ' · '.join(subtitle_parts)
        return (
            f'<header><h1>{self._report_title(s)}</h1>\n'
            + (f'<p class="report-subtitle">{subtitle}</p>\n' if subtitle else '')
            + f'<p class="report-generated">{self._t("generated")}: {generated}</p>\n</header>\n'
        )

    @staticmethod
    def _compliance_score(s: dict[str, Any]) -> float:
        passes = int(s.get('total_passes', 0) or 0)
        violations = int(s.get('total_violations', 0) or 0)
        warnings_count = int(s.get('total_warnings', 0) or 0)
        total = passes + violations + warnings_count
        return (passes / total * 100.0) if total > 0 else 0.0

    @staticmethod
    def _average_page_score(s: dict[str, Any]) -> float | None:
        raw = s.get('page_scores') or []
        values: list[float] = []
        for entry in raw:
            score = entry[1] if isinstance(entry, (tuple, list)) and len(entry) > 1 else None
            if isinstance(score, (int, float)):
                values.append(float(score))
        if not values:
            return None
        avg = sum(values) / len(values)
        # DB scores may be stored as a 0-1 fraction; normalize to percent
        return avg * 100.0 if avg <= 1.0 else avg

    @staticmethod
    def _score_color(score: float) -> str:
        """Ring colour for the compliance circle (non-text UI, 3:1+ on white)."""
        if score >= 90:
            return '#1e8449'
        if score >= 70:
            return '#b38600'
        return '#922b21'

    @staticmethod
    def _grade(score: float) -> str:
        if score >= 90:
            return 'A'
        if score >= 80:
            return 'B'
        if score >= 70:
            return 'C'
        if score >= 60:
            return 'D'
        return 'F'

    def _render_executive_summary(self, s: dict[str, Any]) -> str:
        compliance = self._compliance_score(s)
        colour = self._score_color(compliance)
        dash = compliance * 5.65
        summary_text = self._t('exec_summary_text').format(
            pages=int(s.get('total_pages', 0) or 0),
            violations=int(s.get('total_violations', 0) or 0),
            warnings=int(s.get('total_warnings', 0) or 0),
            passes=int(s.get('total_passes', 0) or 0),
        )
        return f'''<section class="executive-summary">
<h2>{self._t('executive_summary')}</h2>
<div class="exec-grid">
<div class="compliance-score-card">
<div class="score-circle" role="img" aria-label="{self._t('compliance')}: {compliance:.1f}% — {self._t('grade')} {self._grade(compliance)}">
<svg viewBox="0 0 200 200" width="180" height="180" aria-hidden="true">
<circle cx="100" cy="100" r="90" fill="none" stroke="#e9ecef" stroke-width="20"/>
<circle cx="100" cy="100" r="90" fill="none" stroke="{colour}" stroke-width="20" stroke-dasharray="{dash:.1f} 565" transform="rotate(-90 100 100)"/>
</svg>
<div class="score-text"><span class="score-value">{compliance:.1f}%</span><span class="score-label">{self._t('compliance')}</span></div>
</div>
<div class="score-grade">{self._t('grade')}: {self._grade(compliance)}</div>
<p class="score-desc">{self._t('compliance_score_desc')}</p>
</div>
<div class="summary-text"><p>{summary_text}</p></div>
</div>
</section>
'''

    def _render_key_metrics(self, s: dict[str, Any]) -> str:
        avg_score = self._average_page_score(s)
        cards: list[tuple[str, str, str]] = [
            (str(int(s.get('total_pages', 0) or 0)), self._t('total_pages'), ''),
            (str(int(s.get('total_violations', 0) or 0)), self._t('total_violations'), ' violations'),
            (str(int(s.get('total_warnings', 0) or 0)), self._t('total_warnings'), ' warnings'),
            (str(int(s.get('total_info', 0) or 0)), self._t('info'), ''),
            (str(int(s.get('total_discovery', 0) or 0)), self._t('discovery'), ''),
            (str(int(s.get('total_passes', 0) or 0)), self._t('passes'), ' passes'),
        ]
        if avg_score is not None:
            cards.append((f'{avg_score:.0f}%', self._t('average_page_score'), ''))
        html_out = f'<section class="key-metrics"><h2>{self._t("key_metrics")}</h2><div class="stats-grid">\n'
        for value, label, css_class in cards:
            html_out += f'<div class="stat-card{css_class}"><h3>{value}</h3><p>{label}</p></div>\n'
        html_out += '</div></section>\n'
        return html_out

    def _touchpoint_label(self, tp_id: str) -> str:
        """Human-readable, translated touchpoint name."""
        key = f'tp_{tp_id}'
        translated = self._t(key)
        if translated != key:
            return translated
        try:
            from auto_a11y.core.touchpoints import TouchpointID, get_touchpoint
            tp = get_touchpoint(TouchpointID(tp_id))
            if tp is not None:
                return tp.name
        except ValueError:
            pass
        return tp_id.replace('_', ' ').title()

    def _render_bar_rows(self, rows: list[tuple[str, int, str]]) -> str:
        """Render label/count bar rows. Value text sits outside the bar."""
        if not rows:
            return ''
        max_count = max(count for _, count, _ in rows) or 1
        out = '<ul class="bar-list">\n'
        for label, count, colour in rows:
            pct = max(count / max_count * 100.0, 1.5)
            out += (
                f'<li><span class="bar-label">{self._esc(label)}</span>'
                f'<span class="bar-track"><span class="bar" style="width:{pct:.1f}%;background-color:{colour}"></span></span>'
                f'<span class="bar-value">{count}</span></li>\n'
            )
        out += '</ul>\n'
        return out

    def _render_touchpoint_chart(self, s: dict[str, Any]) -> str:
        counts_raw = s.get('touchpoint_counts') or {}
        counts: list[tuple[str, int]] = sorted(
            ((str(k), int(v)) for k, v in counts_raw.items() if int(v) > 0),
            key=lambda kv: kv[1], reverse=True
        )[:12]
        if not counts:
            return ''
        rows = [(self._touchpoint_label(tp), count, '#0a58ca') for tp, count in counts]
        return (
            f'<section class="chart-section"><h2>{self._t("issues_by_touchpoint")}</h2>\n'
            + self._render_bar_rows(rows) + '</section>\n'
        )

    def _render_impact_chart(self, s: dict[str, Any]) -> str:
        counts_raw = s.get('impact_counts') or {}
        impact_colours = {'high': '#922b21', 'medium': '#7d6608', 'low': '#2c3e50'}
        rows: list[tuple[str, int, str]] = []
        for impact in ('high', 'medium', 'low'):
            count = int(counts_raw.get(impact, 0) or 0)
            if count > 0:
                rows.append((self._translate_impact(impact), count, impact_colours[impact]))
        if not rows:
            return ''
        return (
            f'<section class="chart-section"><h2>{self._t("issues_by_impact")}</h2>\n'
            + self._render_bar_rows(rows) + '</section>\n'
        )

    def _render_top_issues(self, s: dict[str, Any]) -> str:
        top_raw = s.get('top_issue_codes') or {}
        items: list[tuple[str, int]] = sorted(
            ((str(k), int(v)) for k, v in dict(top_raw).items()),
            key=lambda kv: kv[1], reverse=True
        )[:10]
        if not items:
            return ''
        out = (
            f'<section class="top-issues"><h2>{self._t("top_issues")}</h2>\n'
            f'<table><thead><tr><th>{self._t("code")}</th><th>{self._t("description")}</th>'
            f'<th>{self._t("occurrences")}</th></tr></thead><tbody>\n'
        )
        for code, count in items:
            title = IssueCatalog.get_issue(code).get('title', '') or code
            out += f'<tr><td>{self._esc(code)}</td><td>{self._esc(title)}</td><td>{count}</td></tr>\n'
        out += '</tbody></table></section>\n'
        return out

    def _render_recommendations(self, s: dict[str, Any]) -> str:
        recs: list[str] = []
        impact_raw = s.get('impact_counts') or {}
        high_count = int(impact_raw.get('high', 0) or 0)
        violations = int(s.get('total_violations', 0) or 0)
        warnings_count = int(s.get('total_warnings', 0) or 0)

        if violations == 0:
            recs.append(self._t('rec_none'))
        else:
            if high_count > 0:
                recs.append(self._t('rec_high_impact').format(count=high_count))
            counts_raw = s.get('touchpoint_counts') or {}
            top_tp = sorted(
                ((str(k), int(v)) for k, v in counts_raw.items() if int(v) > 0),
                key=lambda kv: kv[1], reverse=True
            )[:1]
            for tp, count in top_tp:
                recs.append(self._t('rec_touchpoint').format(name=self._touchpoint_label(tp), count=count))
        if warnings_count > 0:
            recs.append(self._t('rec_warnings').format(count=warnings_count))
        recs.append(self._t('rec_retest'))

        out = f'<section class="recommendations"><h2>{self._t("recommendations")}</h2><ol>\n'
        for rec in recs:
            out += f'<li>{rec}</li>\n'
        out += '</ol></section>\n'
        return out

    @staticmethod
    def _get_exec_css() -> str:
        """Styles for the executive sections. HTML streaming report only."""
        return """
    <style>
        .report-subtitle { color: #555; font-size: 1.05em; }
        .report-generated { color: #555; font-size: 0.9em; }
        .exec-grid {
            display: grid;
            grid-template-columns: auto 1fr;
            gap: 30px;
            align-items: center;
        }
        @media (max-width: 700px) { .exec-grid { grid-template-columns: 1fr; } }
        .compliance-score-card { text-align: center; }
        .score-circle { position: relative; display: inline-block; }
        .score-text {
            position: absolute;
            top: 50%; left: 50%;
            transform: translate(-50%, -50%);
            display: flex; flex-direction: column;
        }
        .score-value { font-size: 1.8em; font-weight: 700; color: #1c1c1c; }
        .score-label { font-size: 0.9em; color: #555; }
        .score-grade { font-size: 1.2em; font-weight: 700; margin-top: 6px; }
        .score-desc { color: #555; font-size: 0.9em; max-width: 240px; margin: 6px auto 0; }
        .summary-text p { font-size: 1.1em; }
        .bar-list { list-style: none; margin: 15px 0; padding: 0; }
        .bar-list li {
            display: grid;
            grid-template-columns: minmax(140px, 220px) 1fr 3.5em;
            gap: 10px;
            align-items: center;
            margin-bottom: 8px;
        }
        .bar-label { font-size: 0.95em; overflow-wrap: anywhere; }
        .bar-track {
            display: block;
            background: #e9ecef;
            border-radius: 4px;
            height: 20px;
            overflow: hidden;
        }
        .bar { display: block; height: 100%; border-radius: 4px; }
        .bar-value { font-variant-numeric: tabular-nums; text-align: right; font-weight: 600; }
        .recommendations ol { margin: 10px 0 10px 25px; }
        .recommendations li { margin-bottom: 8px; }
        .detailed-results { margin-top: 30px; }
    </style>
"""


class JSONFormatter(BaseFormatter):
    """JSON report formatter"""

    _file: TextIO | None
    _is_first_page: bool

    def __init__(self, config: dict[str, Any], language: str = 'en') -> None:
        super().__init__(config, language)
        self.extension = 'json'
        self._file = None
        self._is_first_page = True

    @override
    def format_page_report(self, data: dict[str, Any]) -> str:
        """Generate JSON report for a page"""
        return json.dumps(data, indent=2, default=str)

    @override
    def format_website_report(self, data: dict[str, Any]) -> str:
        """Generate JSON report for a website"""
        return json.dumps(data, indent=2, default=str)

    @override
    def format_project_report(self, data: dict[str, Any]) -> str:
        """Generate JSON report for a project"""
        return json.dumps(data, indent=2, default=str)

    @override
    def format_summary_report(self, data: dict[str, Any]) -> str:
        """Generate JSON summary report"""
        return json.dumps(data, indent=2, default=str)

    # --- Streaming interface ---

    @override
    def begin(self, output_file: str, summary: dict[str, Any]) -> None:
        """Write the opening envelope: {"summary": ..., "pages": ["""
        self._file = open(output_file, 'w', encoding='utf-8')
        self._is_first_page = True
        self._file.write('{"summary": ')
        json.dump(summary, self._file, indent=2, default=str)
        self._file.write(', "pages": [\n')

    @override
    def append_page(self, output_file: str, page_data: dict[str, Any]) -> None:
        """Write one page JSON object into the pages array."""
        if self._file is None:
            return
        if not self._is_first_page:
            self._file.write(',\n')
        self._is_first_page = False

        # Serialise page_data; convert model objects to dicts where needed
        serialisable = self._make_serialisable(page_data)
        json.dump(serialisable, self._file, indent=2, default=str)

    @override
    def finalize(self, output_file: str, summary: dict[str, Any]) -> None:
        """Close the pages array and the root object."""
        if self._file is not None and not self._file.closed:
            self._file.write('\n]}')
            self._file.flush()
            self._file.close()

    @override
    def cleanup(self) -> None:
        """Close the file handle if still open."""
        if self._file is not None and not self._file.closed:
            self._file.close()

    # --- helpers ---

    @staticmethod
    def _make_serialisable(obj: object) -> object:
        """Recursively convert model objects (with __dict__) to plain dicts."""
        if hasattr(obj, 'items') and hasattr(obj, 'keys'):
            # dict-like -- extract items via attribute protocol
            items: list[tuple[str, object]] = [(str(k), v) for k, v in getattr(obj, 'items')()]
            return {k: JSONFormatter._make_serialisable(v) for k, v in items}
        if hasattr(obj, '__iter__') and not isinstance(obj, (str, bytes)):
            # list/tuple-like -- iterate via iter()
            elems: list[object] = list(getattr(obj, '__iter__')())
            return [JSONFormatter._make_serialisable(e) for e in elems]
        if hasattr(obj, '__dict__') and not isinstance(obj, type):
            attrs: dict[str, object] = dict(vars(obj))
            return {k: JSONFormatter._make_serialisable(v)
                    for k, v in attrs.items() if not k.startswith('_')}
        if hasattr(obj, 'value'):  # enum-like
            return getattr(obj, 'value')
        return obj


class CSVFormatter(BaseFormatter):
    """CSV report formatter"""

    _file: TextIO | None

    def __init__(self, config: dict[str, Any], language: str = 'en') -> None:
        super().__init__(config, language)
        self.extension = 'csv'
        self._file: TextIO | None = None
        self._writer = csv.writer(StringIO())  # placeholder; replaced in begin()

    @override
    def format_page_report(self, data: dict[str, Any]) -> str:
        """Generate CSV report for a page"""
        output = StringIO()
        writer = csv.writer(output)

        # Write header
        writer.writerow([self._t('type'), self._t('rule_id'), self._t('description'), self._t('impact'), self._t('wcag_criteria'), self._t('elements_affected'), self._t('fix')])

        # Write violations
        for v in data['violations']:
            writer.writerow([
                self._t('violation'),
                v.get('rule_id', ''),
                v.get('description', ''),
                v.get('impact', ''),
                ', '.join(v.get('wcag_criteria', [])),
                v.get('node_count', 0),
                v.get('suggested_fix', '')
            ])

        # Write warnings
        for w in data['warnings']:
            writer.writerow([
                self._t('warning'),
                w.get('rule_id', ''),
                w.get('description', ''),
                '',
                '',
                w.get('node_count', 0),
                ''
            ])

        return output.getvalue()

    @override
    def format_website_report(self, data: dict[str, Any]) -> str:
        """Generate CSV report for a website"""
        output = StringIO()
        writer = csv.writer(output)

        # Write header
        writer.writerow([self._t('page_url'), self._t('violations'), self._t('warnings'), self._t('passes'), self._t('last_tested')])
        
        # Write page data
        for pr in data['pages']:
            page = pr['page']
            test = pr['test_result']
            writer.writerow([
                page.url,
                test.violation_count,
                test.warning_count,
                test.pass_count,
                test.test_date
            ])
        
        return output.getvalue()
    
    @override
    def format_project_report(self, data: dict[str, Any]) -> str:
        """Generate CSV report for a project"""
        output = StringIO()
        writer = csv.writer(output)
        
        # Write header
        writer.writerow([self._t('website'), self._t('url'), self._t('pages'), self._t('total_violations'), self._t('total_warnings')])
        
        # Write website data
        for wd in data['websites']:
            website = wd['website']
            pages = wd['pages']
            
            total_violations = sum(p['test_result'].violation_count for p in pages)
            total_warnings = sum(p['test_result'].warning_count for p in pages)
            
            writer.writerow([
                website.name,
                website.url,
                len(pages),
                total_violations,
                total_warnings
            ])
        
        return output.getvalue()
    
    @override
    def format_summary_report(self, data: dict[str, Any]) -> str:
        """Generate CSV summary report"""
        output = StringIO()
        writer = csv.writer(output)
        
        # Write header
        writer.writerow([self._t('project'), self._t('websites'), self._t('pages_tested'), self._t('violations')])
        
        # Write project data
        for project in data['projects']:
            writer.writerow([
                project['name'],
                project['websites'],
                project['pages_tested'],
                project['violations']
            ])
        
        return output.getvalue()

    # --- Streaming interface ---

    def _csv_header(self) -> list[str]:
        """Get translated CSV header row"""
        return [self._t('url'), self._t('page_title'), self._t('type'), self._t('code'), self._t('description'),
                self._t('touchpoint'), self._t('impact'), self._t('xpath'), self._t('html'), self._t('wcag_criteria')]

    @override
    def begin(self, output_file: str, summary: dict[str, Any]) -> None:
        """Open *output_file* and write the CSV header row."""
        self._file = open(output_file, 'w', newline='', encoding='utf-8')
        self._writer = csv.writer(self._file)
        self._writer.writerow(self._csv_header())

    @override
    def append_page(self, output_file: str, page_data: dict[str, Any]) -> None:
        """Write rows for every violation and warning on one page."""
        page = page_data.get('page', {})
        test_result = page_data.get('test_result')
        if test_result is None:
            return

        page_url: str = self.get_page_url(page)
        page_title: str = self.get_page_title(page)

        violations: list[Any] = self.get_issue_list(test_result, 'violations')
        warnings_list: list[Any] = self.get_issue_list(test_result, 'warnings')

        for v in violations:
            self._write_issue_row('Error', page_url, page_title, v)
        for w in warnings_list:
            self._write_issue_row('Warning', page_url, page_title, w)

    @override
    def finalize(self, output_file: str, summary: dict[str, Any]) -> None:
        """Flush and close the CSV file."""
        if self._file is not None and not self._file.closed:
            self._file.flush()
            self._file.close()

    @override
    def cleanup(self) -> None:
        """Close the file handle if still open."""
        if self._file is not None and not self._file.closed:
            self._file.close()

    # --- helpers ---

    def _write_issue_row(self, issue_type: str, page_url: str, page_title: str, issue: Any) -> None:
        """Write a single CSV row for an issue (violation or warning)."""
        issue_dict: dict[str, Any] = self.to_enriched_dict(issue)

        impact_raw = issue_dict.get('impact', '')
        if hasattr(impact_raw, 'value'):
            impact_raw = impact_raw.value

        wcag_str: str = self._format_wcag(issue_dict)

        self._writer.writerow([
            page_url,
            page_title,
            issue_type,
            str(issue_dict.get('id', '')),
            self._best_description(issue_dict),
            str(issue_dict.get('touchpoint', '')),
            self._translate_impact(str(impact_raw)),
            issue_dict.get('xpath', ''),
            issue_dict.get('html', ''),
            wcag_str,
        ])


class ExcelFormatter(BaseFormatter):
    """Excel report formatter - inherits TRANSLATIONS and _t() from BaseFormatter"""

    def __init__(self, config: dict[str, Any], language: str = 'en') -> None:
        super().__init__(config, language)
        self.extension = 'xlsx'
        # Streaming state
        self._output_file: str = ''
        self._wb: Any = None  # openpyxl Workbook
        self._ws_violations: Any = None  # openpyxl Worksheet
        self._ws_warnings: Any = None  # openpyxl Worksheet
        from openpyxl import Workbook
        from openpyxl.styles import Font, Fill, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        self.Workbook = Workbook
        self.Font = Font
        self.Fill = Fill
        self.PatternFill = PatternFill
        self.Alignment = Alignment
        self.Border = Border
        self.Side = Side
        self.get_column_letter = get_column_letter

    def _get_styles(self) -> dict[str, Any]:
        """Get common Excel styles"""
        return {
            'header': {
                'font': self.Font(bold=True, color="FFFFFF", size=12),
                'fill': self.PatternFill(start_color="007BFF", end_color="007BFF", fill_type="solid"),
                'alignment': self.Alignment(horizontal="center", vertical="center"),
                'border': self.Border(
                    left=self.Side(style='thin'),
                    right=self.Side(style='thin'),
                    top=self.Side(style='thin'),
                    bottom=self.Side(style='thin')
                )
            },
            'subheader': {
                'font': self.Font(bold=True, size=11),
                'fill': self.PatternFill(start_color="E9ECEF", end_color="E9ECEF", fill_type="solid"),
                'alignment': self.Alignment(horizontal="left", vertical="center")
            },
            'violation': {
                'fill': self.PatternFill(start_color="FFE5E5", end_color="FFE5E5", fill_type="solid")
            },
            'warning': {
                'fill': self.PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
            },
            'info': {
                'fill': self.PatternFill(start_color="D1ECF1", end_color="D1ECF1", fill_type="solid")
            },
            'discovery': {
                'fill': self.PatternFill(start_color="E6E0FF", end_color="E6E0FF", fill_type="solid")
            },
            'pass': {
                'fill': self.PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
            },
            'critical': {
                'font': self.Font(color="DC3545", bold=True)
            },
            'serious': {
                'font': self.Font(color="FD7E14", bold=True)
            },
            'moderate': {
                'font': self.Font(color="FFC107")
            },
            'minor': {
                'font': self.Font(color="6C757D")
            }
        }
    
    @override
    def format_page_report(self, data: dict[str, Any]) -> bytes:
        """Generate Excel report for a page"""
        wb = self.Workbook()
        styles = self._get_styles()

        # Summary Sheet
        ws_summary = wb.active
        assert ws_summary is not None
        ws_summary.title = self._t('summary')
        self._create_summary_sheet(ws_summary, data, styles)

        # Multi-State Summary Sheet (if applicable)
        test_result = data.get('test_result', {})
        if test_result and test_result.get('session_id'):
            ws_states = wb.create_sheet(self._t('page_states'))
            self._create_page_states_sheet(ws_states, data, styles)

        # All Issues Sheet (combined view)
        ws_all_issues = wb.create_sheet(self._t('all_issues'))
        self._create_all_issues_sheet(ws_all_issues, data, styles)

        # Violations Sheet
        if data.get('violations'):
            ws_violations = wb.create_sheet(self._t('violations'))
            self._create_violations_sheet(ws_violations, data['violations'], styles)
        
        # Warnings Sheet
        if data.get('warnings'):
            ws_warnings = wb.create_sheet(self._t('warnings'))
            self._create_warnings_sheet(ws_warnings, data['warnings'], styles)
        
        # Info Sheet
        if data.get('info'):
            ws_info = wb.create_sheet(self._t('info'))
            self._create_info_sheet(ws_info, data['info'], styles)
        
        # Discovery Sheet
        if data.get('discovery'):
            ws_discovery = wb.create_sheet(self._t('discovery'))
            self._create_discovery_sheet(ws_discovery, data['discovery'], styles)
        
        # AI Findings Sheet
        if data.get('ai_findings'):
            ws_ai = wb.create_sheet(self._t('ai_findings'))
            self._create_ai_findings_sheet(ws_ai, data['ai_findings'], styles)
        
        # Passes Sheet
        if data.get('passes'):
            ws_passes = wb.create_sheet(self._t('passes'))
            self._create_passes_sheet(ws_passes, data['passes'], styles)
        
        # Save to bytes
        from io import BytesIO
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()
    
    @override
    def format_website_report(self, data: dict[str, Any]) -> bytes:
        """Generate Excel report for a website"""
        wb = self.Workbook()
        styles = self._get_styles()
        
        # Summary Sheet
        ws_summary = wb.active
        assert ws_summary is not None
        ws_summary.title = self._t('website_summary')
        self._create_website_summary_sheet(ws_summary, data, styles)
        
        # Pages Sheet
        if data.get('pages'):
            ws_pages = wb.create_sheet(self._t('pages'))
            self._create_pages_sheet(ws_pages, data['pages'], styles)
        
        # Violation Types Sheet
        if data.get('violation_types'):
            ws_types = wb.create_sheet(self._t('violation_types'))
            self._create_violation_types_sheet(ws_types, data['violation_types'], styles)
        
        # Save to bytes
        from io import BytesIO
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()
    
    @override
    def format_project_report(self, data: dict[str, Any]) -> bytes:
        """Generate Excel report for a project"""
        wb = self.Workbook()
        styles = self._get_styles()
        
        # Project Summary Sheet
        ws_summary = wb.active
        assert ws_summary is not None
        ws_summary.title = self._t('project_summary')
        self._create_project_summary_sheet(ws_summary, data, styles)
        
        # Websites Sheet
        if data.get('websites'):
            ws_websites = wb.create_sheet(self._t('websites'))
            self._create_websites_sheet(ws_websites, data['websites'], styles)

        # All Issues Sheet (combined view across all pages)
        ws_all_issues = wb.create_sheet(self._t('all_issues'))
        self._create_project_all_issues_sheet(ws_all_issues, data, styles)

        # All Issues (Deduplicated) Sheet - groups by common components
        ws_deduped_issues = wb.create_sheet(self._t('all_issues_deduplicated'))
        self._create_project_deduped_issues_sheet(ws_deduped_issues, data, styles)

        # Common Components Sheet - shows all identified common components
        ws_common_components = wb.create_sheet(self._t('common_components'))
        self._create_common_components_sheet(ws_common_components, data, styles)

        # Save to bytes
        from io import BytesIO
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()
    
    def format_all_projects_report(self, data: dict[str, Any]) -> bytes:
        """Generate Excel report for all projects"""
        warnings.warn(
            "format_all_projects_report() is deprecated, use begin/append_page/finalize streaming interface",
            DeprecationWarning,
            stacklevel=2
        )
        wb = self.Workbook()
        styles = self._get_styles()
        
        # Overall Summary Sheet
        ws_summary = wb.active
        assert ws_summary is not None
        ws_summary.title = self._t('overall_summary')
        self._create_all_projects_summary_sheet(ws_summary, data, styles)
        
        # Projects Breakdown Sheet
        if data.get('projects'):
            ws_projects = wb.create_sheet(self._t('projects_breakdown'))
            self._create_projects_breakdown_sheet(ws_projects, data['projects'], styles)
        
        # Save to bytes
        from io import BytesIO
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()
    
    @override
    def format_summary_report(self, data: dict[str, Any]) -> bytes:
        """Generate Excel summary report"""
        wb = self.Workbook()
        styles = self._get_styles()
        
        # Executive Summary Sheet
        ws = wb.active
        assert ws is not None
        ws.title = self._t('executive_summary')
        
        # Headers
        headers = [self._t('project'), self._t('websites'), self._t('pages_tested'), self._t('violations'), self._t('warnings')]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            self._apply_style(cell, styles['header'])
        
        # Data
        row = 2
        for project in data.get('projects', []):
            ws.cell(row=row, column=1, value=project.get('name', ''))
            ws.cell(row=row, column=2, value=project.get('websites', 0))
            ws.cell(row=row, column=3, value=project.get('pages_tested', 0))
            
            violations_cell = ws.cell(row=row, column=4, value=project.get('violations', 0))
            if project.get('violations', 0) > 0:
                violations_cell.fill = styles['violation']['fill']
            
            warnings_cell = ws.cell(row=row, column=5, value=project.get('warnings', 0))
            if project.get('warnings', 0) > 0:
                warnings_cell.fill = styles['warning']['fill']
            
            row += 1
        
        # Auto-adjust column widths
        self._auto_adjust_columns(ws)
        
        # Save to bytes
        from io import BytesIO
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()
    
    def _create_summary_sheet(self, ws: Any, data: dict[str, Any], styles: dict[str, Any]) -> None:
        """Create page summary sheet"""
        # Title
        ws.merge_cells('A1:E1')
        title_cell = ws['A1']
        title_cell.value = f"{self._t('accessibility_report')} - {data.get('page', {}).get('url', self._t('unknown'))}"
        title_cell.font = self.Font(bold=True, size=14)
        title_cell.alignment = self.Alignment(horizontal="center")

        # Metadata
        row = 3
        ws.cell(row=row, column=1, value=f"{self._t('website')}:").font = self.Font(bold=True)
        ws.cell(row=row, column=2, value=data.get('website', {}).get('name', ''))
        row += 1
        ws.cell(row=row, column=1, value=f"{self._t('project')}:").font = self.Font(bold=True)
        ws.cell(row=row, column=2, value=data.get('project', {}).get('name', ''))
        row += 1
        ws.cell(row=row, column=1, value=f"{self._t('generated')}:").font = self.Font(bold=True)
        ws.cell(row=row, column=2, value=data.get('generated_at', ''))

        # Statistics
        row = 7
        ws.cell(row=row, column=1, value=self._t('summary_statistics')).font = self.Font(bold=True, size=12)
        row += 1

        stats = data.get('statistics', {})
        ws.cell(row=row, column=1, value=f"{self._t('violations')}:")
        ws.cell(row=row, column=2, value=stats.get('violations', 0))
        if stats.get('violations', 0) > 0:
            ws.cell(row=row, column=2).fill = styles['violation']['fill']
        row += 1

        ws.cell(row=row, column=1, value=f"{self._t('warnings')}:")
        ws.cell(row=row, column=2, value=stats.get('warnings', 0))
        if stats.get('warnings', 0) > 0:
            ws.cell(row=row, column=2).fill = styles['warning']['fill']
        row += 1

        ws.cell(row=row, column=1, value=f"{self._t('passes')}:")
        ws.cell(row=row, column=2, value=stats.get('passes', 0))
        ws.cell(row=row, column=2).fill = styles['pass']['fill']
        row += 1

        ws.cell(row=row, column=1, value=f"{self._t('test_duration')} (ms):")
        ws.cell(row=row, column=2, value=stats.get('duration_ms', 0))
        
        self._auto_adjust_columns(ws)
    
    def _create_violations_sheet(self, ws: Any, violations: list[dict[str, Any]], styles: dict[str, Any]) -> None:
        """Create violations sheet"""
        headers = [self._t('rule_id'), self._t('description'), self._t('impact'), self._t('wcag_criteria'), self._t('elements_affected'), self._t('suggested_fix'), self._t('test_user'), self._t('user_roles')]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            self._apply_style(cell, styles['header'])

        row = 2
        for v in violations:
            ws.cell(row=row, column=1, value=v.get('rule_id', ''))
            ws.cell(row=row, column=2, value=v.get('description', ''))

            impact = v.get('impact', 'moderate')
            impact_cell = ws.cell(row=row, column=3, value=impact.upper())
            if impact.lower() in styles:
                impact_cell.font = styles[impact.lower()]['font']

            ws.cell(row=row, column=4, value=', '.join(v.get('wcag_criteria', [])))
            ws.cell(row=row, column=5, value=v.get('node_count', 0))
            ws.cell(row=row, column=6, value=v.get('suggested_fix', ''))

            # Add authenticated user info
            metadata = v.get('metadata', {})
            auth_user = metadata.get('authenticated_user', {})
            if auth_user:
                ws.cell(row=row, column=7, value=auth_user.get('display_name', ''))
                ws.cell(row=row, column=8, value=', '.join(auth_user.get('roles', [])))
            else:
                ws.cell(row=row, column=7, value='Guest')
                ws.cell(row=row, column=8, value='no login')

            # Apply row coloring
            for col in range(1, 9):
                ws.cell(row=row, column=col).fill = styles['violation']['fill']

            row += 1

        self._auto_adjust_columns(ws)
    
    def _create_warnings_sheet(self, ws: Any, warnings: list[dict[str, Any]], styles: dict[str, Any]) -> None:
        """Create warnings sheet"""
        headers = [self._t('rule_id'), self._t('description'), self._t('elements_affected'), self._t('test_user'), self._t('user_roles')]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            self._apply_style(cell, styles['header'])

        row = 2
        for w in warnings:
            ws.cell(row=row, column=1, value=w.get('rule_id', ''))
            ws.cell(row=row, column=2, value=w.get('description', ''))
            ws.cell(row=row, column=3, value=w.get('node_count', 0))

            # Add authenticated user info
            metadata = w.get('metadata', {})
            auth_user = metadata.get('authenticated_user', {})
            if auth_user:
                ws.cell(row=row, column=4, value=auth_user.get('display_name', ''))
                ws.cell(row=row, column=5, value=', '.join(auth_user.get('roles', [])))
            else:
                ws.cell(row=row, column=4, value='Guest')
                ws.cell(row=row, column=5, value='no login')

            # Apply row coloring
            for col in range(1, 6):
                ws.cell(row=row, column=col).fill = styles['warning']['fill']

            row += 1

        self._auto_adjust_columns(ws)
    
    def _create_info_sheet(self, ws: Any, info_items: list[dict[str, Any]], styles: dict[str, Any]) -> None:
        """Create info sheet"""
        headers = [self._t('id'), self._t('description'), self._t('category'), self._t('wcag_criteria'), self._t('location'), self._t('test_user'), self._t('user_roles')]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            self._apply_style(cell, styles['header'])

        row = 2
        for item in info_items:
            ws.cell(row=row, column=1, value=item.get('id', ''))
            ws.cell(row=row, column=2, value=item.get('description', ''))
            ws.cell(row=row, column=3, value=item.get('category', ''))
            ws.cell(row=row, column=4, value=', '.join(item.get('wcag_criteria', [])))
            ws.cell(row=row, column=5, value=item.get('xpath', ''))

            # Add authenticated user info
            metadata = item.get('metadata', {})
            auth_user = metadata.get('authenticated_user', {})
            if auth_user:
                ws.cell(row=row, column=6, value=auth_user.get('display_name', ''))
                ws.cell(row=row, column=7, value=', '.join(auth_user.get('roles', [])))
            else:
                ws.cell(row=row, column=6, value='Guest')
                ws.cell(row=row, column=7, value='no login')

            # Apply info coloring (blue)
            for col in range(1, 8):
                ws.cell(row=row, column=col).fill = styles['info']['fill']

            row += 1

        self._auto_adjust_columns(ws)
    
    def _create_discovery_sheet(self, ws: Any, discovery_items: list[dict[str, Any]], styles: dict[str, Any]) -> None:
        """Create discovery sheet"""
        headers = [self._t('id'), self._t('description'), self._t('category'), self._t('location'), self._t('manual_check_required'), self._t('test_user'), self._t('user_roles')]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            self._apply_style(cell, styles['header'])

        row = 2
        for item in discovery_items:
            ws.cell(row=row, column=1, value=item.get('id', ''))
            ws.cell(row=row, column=2, value=item.get('description', ''))
            ws.cell(row=row, column=3, value=item.get('category', ''))
            ws.cell(row=row, column=4, value=item.get('xpath', ''))
            ws.cell(row=row, column=5, value='Yes')

            # Add authenticated user info
            metadata = item.get('metadata', {})
            auth_user = metadata.get('authenticated_user', {})
            if auth_user:
                ws.cell(row=row, column=6, value=auth_user.get('display_name', ''))
                ws.cell(row=row, column=7, value=', '.join(auth_user.get('roles', [])))
            else:
                ws.cell(row=row, column=6, value='Guest')
                ws.cell(row=row, column=7, value='no login')

            # Apply discovery coloring (purple)
            for col in range(1, 8):
                cell = ws.cell(row=row, column=col)
                if 'discovery' in styles:
                    cell.fill = styles['discovery']['fill']
                else:
                    # Use a purple-ish color if not defined
                    from openpyxl.styles import PatternFill
                    cell.fill = PatternFill(start_color="E6E0FF", end_color="E6E0FF", fill_type="solid")

            row += 1

        self._auto_adjust_columns(ws)
    
    def _create_ai_findings_sheet(self, ws: Any, findings: list[Any], styles: dict[str, Any]) -> None:
        """Create AI findings sheet"""
        headers = [self._t('type'), self._t('description'), self._t('severity'), self._t('confidence'), self._t('suggested_fix')]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            self._apply_style(cell, styles['header'])
        
        row = 2
        for f in findings:
            ws.cell(row=row, column=1, value=getattr(f, 'type', 'AI Finding'))
            ws.cell(row=row, column=2, value=getattr(f, 'description', ''))
            
            severity = getattr(f, 'severity', None)
            if severity:
                severity_value = severity.value if hasattr(severity, 'value') else str(severity)
                severity_cell = ws.cell(row=row, column=3, value=severity_value.upper())
                if severity_value.lower() in styles:
                    severity_cell.font = styles[severity_value.lower()]['font']
            
            confidence = getattr(f, 'confidence', 0.85)
            ws.cell(row=row, column=4, value=f"{confidence * 100:.0f}%")
            ws.cell(row=row, column=5, value=getattr(f, 'suggested_fix', ''))
            
            row += 1
        
        self._auto_adjust_columns(ws)
    
    def _create_passes_sheet(self, ws: Any, passes: list[dict[str, Any]], styles: dict[str, Any]) -> None:
        """Create passes sheet"""
        headers = [self._t('rule_id'), self._t('description')]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            self._apply_style(cell, styles['header'])
        
        row = 2
        for p in passes:
            ws.cell(row=row, column=1, value=p.get('rule_id', ''))
            ws.cell(row=row, column=2, value=p.get('description', 'Passed'))
            
            # Apply row coloring
            for col in range(1, 3):
                ws.cell(row=row, column=col).fill = styles['pass']['fill']
            
            row += 1
        
        self._auto_adjust_columns(ws)

    def _create_all_issues_sheet(self, ws: Any, data: dict[str, Any], styles: dict[str, Any]) -> None:
        """Create a combined sheet with all issues (violations, warnings, info, discovery)"""
        headers = [self._t('type'), self._t('impact'), self._t('rule_id'), self._t('touchpoint'), self._t('what'), self._t('why_important'), self._t('who_affected'), self._t('how_to_remediate'), self._t('wcag_criteria'), self._t('location_xpath'), self._t('element'), self._t('page_url'), self._t('breakpoint_px'), self._t('pseudoclass'), self._t('page_state'), self._t('test_user'), self._t('user_roles')]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            self._apply_style(cell, styles['header'])

        row = 2

        # Get page state description from test result if available
        page_state_desc: str = ''
        test_result = data.get('test_result', {})
        if test_result:
            page_state = test_result.get('page_state')
            if page_state:
                page_state_desc = str(getattr(page_state, 'description', '') if hasattr(page_state, 'description') else (page_state.get('description', '') if hasattr(page_state, 'get') else ''))

        # Add all violations
        for v in data.get('violations', []):
            ws.cell(row=row, column=1, value='Violation')
            ws.cell(row=row, column=2, value=v.get('impact', 'Unknown').upper())
            ws.cell(row=row, column=3, value=v.get('id', v.get('rule_id', '')))
            ws.cell(row=row, column=4, value=v.get('touchpoint', v.get('category', '')))
            ws.cell(row=row, column=5, value=v.get('description_full', v.get('what', v.get('description', ''))))
            ws.cell(row=row, column=6, value=v.get('why_it_matters', ''))
            ws.cell(row=row, column=7, value=v.get('who_it_affects', ''))
            ws.cell(row=row, column=8, value=v.get('how_to_fix', v.get('remediation', v.get('suggested_fix', ''))))
            ws.cell(row=row, column=9, value=v.get('wcag_full', ', '.join(v.get('wcag_criteria', [])) if isinstance(v.get('wcag_criteria'), list) else v.get('wcag_criteria', '')))
            ws.cell(row=row, column=10, value=v.get('xpath', ''))
            ws.cell(row=row, column=11, value=v.get('element', ''))
            ws.cell(row=row, column=12, value=v.get('url', ''))

            # Add metadata columns
            metadata = v.get('metadata', {})
            ws.cell(row=row, column=13, value=metadata.get('breakpoint', ''))
            ws.cell(row=row, column=14, value=metadata.get('pseudoclass', ''))
            ws.cell(row=row, column=15, value=page_state_desc)

            # Add authenticated user info
            auth_user = metadata.get('authenticated_user', {})
            if auth_user:
                ws.cell(row=row, column=16, value=auth_user.get('display_name', ''))
                ws.cell(row=row, column=17, value=', '.join(auth_user.get('roles', [])))
            else:
                ws.cell(row=row, column=16, value='Guest')
                ws.cell(row=row, column=17, value='no login')

            # Apply violation coloring
            for col in range(1, 18):
                ws.cell(row=row, column=col).fill = styles['violation']['fill']

            row += 1

        # Add all warnings
        for w in data.get('warnings', []):
            ws.cell(row=row, column=1, value='Warning')
            ws.cell(row=row, column=2, value=w.get('impact', 'MEDIUM').upper())
            ws.cell(row=row, column=3, value=w.get('id', w.get('rule_id', '')))
            ws.cell(row=row, column=4, value=w.get('touchpoint', w.get('category', '')))
            ws.cell(row=row, column=5, value=w.get('description_full', w.get('what', w.get('description', ''))))
            ws.cell(row=row, column=6, value=w.get('why_it_matters', ''))
            ws.cell(row=row, column=7, value=w.get('who_it_affects', ''))
            ws.cell(row=row, column=8, value=w.get('how_to_fix', w.get('remediation', w.get('suggested_fix', ''))))
            ws.cell(row=row, column=9, value=w.get('wcag_full', ', '.join(w.get('wcag_criteria', [])) if isinstance(w.get('wcag_criteria'), list) else w.get('wcag_criteria', '')))
            ws.cell(row=row, column=10, value=w.get('xpath', ''))
            ws.cell(row=row, column=11, value=w.get('element', ''))
            ws.cell(row=row, column=12, value=w.get('url', ''))

            # Add metadata columns
            metadata = w.get('metadata', {})
            ws.cell(row=row, column=13, value=metadata.get('breakpoint', ''))
            ws.cell(row=row, column=14, value=metadata.get('pseudoclass', ''))
            ws.cell(row=row, column=15, value=page_state_desc)

            # Add authenticated user info
            auth_user = metadata.get('authenticated_user', {})
            if auth_user:
                ws.cell(row=row, column=16, value=auth_user.get('display_name', ''))
                ws.cell(row=row, column=17, value=', '.join(auth_user.get('roles', [])))
            else:
                ws.cell(row=row, column=16, value='Guest')
                ws.cell(row=row, column=17, value='no login')

            # Apply warning coloring
            for col in range(1, 18):
                ws.cell(row=row, column=col).fill = styles['warning']['fill']

            row += 1

        # Add all info items
        for i in data.get('info', []):
            ws.cell(row=row, column=1, value='Info')
            ws.cell(row=row, column=2, value='INFO')
            ws.cell(row=row, column=3, value=i.get('id', ''))
            ws.cell(row=row, column=4, value=i.get('touchpoint', i.get('category', '')))
            ws.cell(row=row, column=5, value=i.get('description_full', i.get('what', i.get('description', ''))))
            ws.cell(row=row, column=6, value=i.get('why_it_matters', ''))
            ws.cell(row=row, column=7, value=i.get('who_it_affects', ''))
            ws.cell(row=row, column=8, value=i.get('how_to_fix', i.get('remediation', '')))
            ws.cell(row=row, column=9, value=i.get('wcag_full', ', '.join(i.get('wcag_criteria', [])) if isinstance(i.get('wcag_criteria'), list) else i.get('wcag_criteria', '')))
            ws.cell(row=row, column=10, value=i.get('xpath', ''))
            ws.cell(row=row, column=11, value=i.get('element', ''))
            ws.cell(row=row, column=12, value=i.get('url', ''))

            # Add metadata columns
            metadata = i.get('metadata', {})
            ws.cell(row=row, column=13, value=metadata.get('breakpoint', ''))
            ws.cell(row=row, column=14, value=metadata.get('pseudoclass', ''))
            ws.cell(row=row, column=15, value=page_state_desc)

            # Add authenticated user info
            auth_user = metadata.get('authenticated_user', {})
            if auth_user:
                ws.cell(row=row, column=16, value=auth_user.get('display_name', ''))
                ws.cell(row=row, column=17, value=', '.join(auth_user.get('roles', [])))
            else:
                ws.cell(row=row, column=16, value='Guest')
                ws.cell(row=row, column=17, value='no login')

            # Apply info coloring
            for col in range(1, 18):
                ws.cell(row=row, column=col).fill = styles['info']['fill']

            row += 1

        # Add all discovery items
        for d in data.get('discovery', []):
            ws.cell(row=row, column=1, value='Discovery')
            ws.cell(row=row, column=2, value='DISCOVERY')
            ws.cell(row=row, column=3, value=d.get('id', d.get('err', '')))
            ws.cell(row=row, column=4, value=d.get('touchpoint', d.get('category', '')))
            ws.cell(row=row, column=5, value=d.get('description_full', d.get('what', d.get('description', ''))))
            ws.cell(row=row, column=6, value=d.get('why_it_matters', ''))
            ws.cell(row=row, column=7, value=d.get('who_it_affects', ''))
            ws.cell(row=row, column=8, value=d.get('how_to_fix', d.get('remediation', '')))
            ws.cell(row=row, column=9, value=d.get('wcag_full', ', '.join(d.get('wcag_criteria', [])) if isinstance(d.get('wcag_criteria'), list) else d.get('wcag_criteria', '')))
            ws.cell(row=row, column=10, value=d.get('xpath', ''))
            ws.cell(row=row, column=11, value=d.get('element', ''))
            ws.cell(row=row, column=12, value=d.get('url', ''))

            # Add metadata columns
            metadata = d.get('metadata', {})
            ws.cell(row=row, column=13, value=metadata.get('breakpoint', ''))
            ws.cell(row=row, column=14, value=metadata.get('pseudoclass', ''))
            ws.cell(row=row, column=15, value=page_state_desc)

            # Add authenticated user info
            auth_user = metadata.get('authenticated_user', {})
            if auth_user:
                ws.cell(row=row, column=16, value=auth_user.get('display_name', ''))
                ws.cell(row=row, column=17, value=', '.join(auth_user.get('roles', [])))
            else:
                ws.cell(row=row, column=16, value='Guest')
                ws.cell(row=row, column=17, value='no login')

            # Apply discovery coloring
            for col in range(1, 18):
                cell = ws.cell(row=row, column=col)
                if 'discovery' in styles:
                    cell.fill = styles['discovery']['fill']
                else:
                    from openpyxl.styles import PatternFill
                    cell.fill = PatternFill(start_color="E6E0FF", end_color="E6E0FF", fill_type="solid")

            row += 1

        # Add AI findings if present
        for f in data.get('ai_findings', []):
            ws.cell(row=row, column=1, value='AI Finding')
            # severity is an ImpactLevel enum (no .upper()); mirror the
            # _create_ai_findings_sheet handling and read .value first.
            sev: object = getattr(f, 'severity', None)
            sev_value: str = sev.value if isinstance(sev, ImpactLevel) else str(sev or 'MEDIUM')
            ws.cell(row=row, column=2, value=sev_value.upper())
            ws.cell(row=row, column=3, value=getattr(f, 'type', ''))
            ws.cell(row=row, column=4, value='')  # AI findings don't have touchpoint
            ws.cell(row=row, column=5, value=getattr(f, 'description', ''))
            ws.cell(row=row, column=6, value='')
            ws.cell(row=row, column=7, value='')
            ws.cell(row=row, column=8, value=getattr(f, 'suggested_fix', ''))
            ws.cell(row=row, column=9, value='')
            ws.cell(row=row, column=10, value='')
            ws.cell(row=row, column=11, value='')
            ws.cell(row=row, column=12, value=getattr(f, 'url', ''))

            # Apply AI finding coloring (use info style)
            for col in range(1, 13):
                ws.cell(row=row, column=col).fill = styles['info']['fill']

            row += 1

        self._auto_adjust_columns(ws)

    def _create_project_all_issues_sheet(self, ws: Any, data: dict[str, Any], styles: dict[str, Any]) -> None:
        """Create a combined sheet with all issues from all pages across all websites"""
        headers = [self._t('type'), self._t('impact'), self._t('rule_id'), self._t('touchpoint'), self._t('what'), self._t('why_important'), self._t('who_affected'), self._t('how_to_remediate'), self._t('wcag_criteria'), self._t('location_xpath'), self._t('element'), self._t('page_url'), self._t('website'), self._t('breakpoint_px'), self._t('pseudoclass'), self._t('page_state'), self._t('test_user'), self._t('user_roles')]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            self._apply_style(cell, styles['header'])

        row = 2

        # Iterate through all websites and their pages
        for website_data in data.get('websites', []):
            website = website_data.get('website', {})
            website_name: str = str(getattr(website, 'name', '') if hasattr(website, 'name') else (website.get('name', '') if hasattr(website, 'get') else ''))

            for page_result in website_data.get('pages', []):
                page = page_result.get('page', {})
                page_url: str = self.get_page_url(page)

                test_result = page_result.get('test_result')
                if not test_result:
                    continue

                # Get page state description from test result if available
                page_state_desc: str = ''
                if test_result:
                    page_state = getattr(test_result, 'page_state', None) if hasattr(test_result, 'page_state') else test_result.get('page_state')
                    if page_state:
                        page_state_desc = str(getattr(page_state, 'description', '') if hasattr(page_state, 'description') else (page_state.get('description', '') if hasattr(page_state, 'get') else ''))

                # Add violations (from list attributes) - enrich with catalog data
                violations = getattr(test_result, 'violations', []) if hasattr(test_result, 'violations') else []
                for v in violations:
                    v_dict: dict[str, Any] = self.to_enriched_dict(v)

                    ws.cell(row=row, column=1, value=self._t('violation'))
                    ws.cell(row=row, column=2, value=self._translate_impact(str(v_dict.get('impact', 'Unknown'))))
                    ws.cell(row=row, column=3, value=v_dict.get('id', ''))
                    ws.cell(row=row, column=4, value=v_dict.get('touchpoint', v_dict.get('category', '')))
                    ws.cell(row=row, column=5, value=self._best_description(v_dict))
                    ws.cell(row=row, column=6, value=v_dict.get('why_it_matters', ''))
                    ws.cell(row=row, column=7, value=v_dict.get('who_it_affects', ''))
                    ws.cell(row=row, column=8, value=v_dict.get('how_to_fix', v_dict.get('remediation', v_dict.get('suggested_fix', ''))))
                    ws.cell(row=row, column=9, value=v_dict.get('wcag_full', ', '.join(v_dict.get('wcag_criteria', [])) if isinstance(v_dict.get('wcag_criteria'), list) else v_dict.get('wcag_criteria', '')))
                    ws.cell(row=row, column=10, value=v_dict.get('xpath', ''))
                    ws.cell(row=row, column=11, value=v_dict.get('element', ''))
                    ws.cell(row=row, column=12, value=page_url)
                    ws.cell(row=row, column=13, value=website_name)

                    # Add metadata columns
                    metadata = v_dict.get('metadata', {})
                    ws.cell(row=row, column=14, value=metadata.get('breakpoint', ''))
                    ws.cell(row=row, column=15, value=metadata.get('pseudoclass', ''))
                    ws.cell(row=row, column=16, value=page_state_desc)

                    # Add authenticated user info
                    auth_user = metadata.get('authenticated_user', {})
                    if auth_user:
                        ws.cell(row=row, column=17, value=auth_user.get('display_name', ''))
                        ws.cell(row=row, column=18, value=', '.join(auth_user.get('roles', [])))
                    else:
                        ws.cell(row=row, column=17, value='Guest')
                        ws.cell(row=row, column=18, value='no login')

                    for col in range(1, 19):
                        ws.cell(row=row, column=col).fill = styles['violation']['fill']
                    row += 1

                # Add warnings - enrich with catalog data
                warnings = getattr(test_result, 'warnings', []) if hasattr(test_result, 'warnings') else []
                for w in warnings:
                    w_dict: dict[str, Any] = self.to_enriched_dict(w)

                    ws.cell(row=row, column=1, value=self._t('warning'))
                    ws.cell(row=row, column=2, value=self._translate_impact(str(w_dict.get('impact', 'Moderate'))))
                    ws.cell(row=row, column=3, value=w_dict.get('id', ''))
                    ws.cell(row=row, column=4, value=w_dict.get('touchpoint', w_dict.get('category', '')))
                    ws.cell(row=row, column=5, value=self._best_description(w_dict))
                    ws.cell(row=row, column=6, value=w_dict.get('why_it_matters', ''))
                    ws.cell(row=row, column=7, value=w_dict.get('who_it_affects', ''))
                    ws.cell(row=row, column=8, value=w_dict.get('how_to_fix', w_dict.get('remediation', w_dict.get('suggested_fix', ''))))
                    ws.cell(row=row, column=9, value=w_dict.get('wcag_full', ', '.join(w_dict.get('wcag_criteria', [])) if isinstance(w_dict.get('wcag_criteria'), list) else w_dict.get('wcag_criteria', '')))
                    ws.cell(row=row, column=10, value=w_dict.get('xpath', ''))
                    ws.cell(row=row, column=11, value=w_dict.get('element', ''))
                    ws.cell(row=row, column=12, value=page_url)
                    ws.cell(row=row, column=13, value=website_name)

                    # Add metadata columns
                    metadata = w_dict.get('metadata', {})
                    ws.cell(row=row, column=14, value=metadata.get('breakpoint', ''))
                    ws.cell(row=row, column=15, value=metadata.get('pseudoclass', ''))
                    ws.cell(row=row, column=16, value=page_state_desc)

                    # Add authenticated user info
                    auth_user = metadata.get('authenticated_user', {})
                    if auth_user:
                        ws.cell(row=row, column=17, value=auth_user.get('display_name', ''))
                        ws.cell(row=row, column=18, value=', '.join(auth_user.get('roles', [])))
                    else:
                        ws.cell(row=row, column=17, value='Guest')
                        ws.cell(row=row, column=18, value='no login')

                    for col in range(1, 19):
                        ws.cell(row=row, column=col).fill = styles['warning']['fill']
                    row += 1

                # Add info items - enrich with catalog data
                info_items = getattr(test_result, 'info', []) if hasattr(test_result, 'info') else []
                for i in info_items:
                    i_dict: dict[str, Any] = self.to_enriched_dict(i)

                    ws.cell(row=row, column=1, value=self._t('info'))
                    ws.cell(row=row, column=2, value='INFO')
                    ws.cell(row=row, column=3, value=i_dict.get('id', ''))
                    ws.cell(row=row, column=4, value=i_dict.get('touchpoint', i_dict.get('category', '')))
                    ws.cell(row=row, column=5, value=self._best_description(i_dict))
                    ws.cell(row=row, column=6, value=i_dict.get('why_it_matters', ''))
                    ws.cell(row=row, column=7, value=i_dict.get('who_it_affects', ''))
                    ws.cell(row=row, column=8, value=i_dict.get('how_to_fix', i_dict.get('remediation', '')))
                    ws.cell(row=row, column=9, value=i_dict.get('wcag_full', ', '.join(i_dict.get('wcag_criteria', [])) if isinstance(i_dict.get('wcag_criteria'), list) else i_dict.get('wcag_criteria', '')))
                    ws.cell(row=row, column=10, value=i_dict.get('xpath', ''))
                    ws.cell(row=row, column=11, value=i_dict.get('element', ''))
                    ws.cell(row=row, column=12, value=page_url)
                    ws.cell(row=row, column=13, value=website_name)

                    # Add metadata columns
                    metadata = i_dict.get('metadata', {})
                    ws.cell(row=row, column=14, value=metadata.get('breakpoint', ''))
                    ws.cell(row=row, column=15, value=metadata.get('pseudoclass', ''))
                    ws.cell(row=row, column=16, value=page_state_desc)

                    # Add authenticated user info
                    auth_user = metadata.get('authenticated_user', {})
                    if auth_user:
                        ws.cell(row=row, column=17, value=auth_user.get('display_name', ''))
                        ws.cell(row=row, column=18, value=', '.join(auth_user.get('roles', [])))
                    else:
                        ws.cell(row=row, column=17, value='Guest')
                        ws.cell(row=row, column=18, value='no login')

                    for col in range(1, 19):
                        ws.cell(row=row, column=col).fill = styles['info']['fill']
                    row += 1

                # Add discovery items - enrich with catalog data
                discovery_items = getattr(test_result, 'discovery', []) if hasattr(test_result, 'discovery') else []
                for d in discovery_items:
                    d_dict: dict[str, Any] = self.to_enriched_dict(d)

                    ws.cell(row=row, column=1, value=self._t('discovery'))
                    ws.cell(row=row, column=2, value='DISCOVERY')
                    ws.cell(row=row, column=3, value=d_dict.get('id', ''))
                    ws.cell(row=row, column=4, value=d_dict.get('touchpoint', d_dict.get('category', '')))
                    ws.cell(row=row, column=5, value=self._best_description(d_dict))
                    ws.cell(row=row, column=6, value=d_dict.get('why_it_matters', ''))
                    ws.cell(row=row, column=7, value=d_dict.get('who_it_affects', ''))
                    ws.cell(row=row, column=8, value=d_dict.get('how_to_fix', d_dict.get('remediation', '')))
                    ws.cell(row=row, column=9, value=d_dict.get('wcag_full', ', '.join(d_dict.get('wcag_criteria', [])) if isinstance(d_dict.get('wcag_criteria'), list) else d_dict.get('wcag_criteria', '')))
                    ws.cell(row=row, column=10, value=d_dict.get('xpath', ''))
                    ws.cell(row=row, column=11, value=d_dict.get('element', ''))
                    ws.cell(row=row, column=12, value=page_url)
                    ws.cell(row=row, column=13, value=website_name)

                    # Add metadata columns
                    metadata = d_dict.get('metadata', {})
                    ws.cell(row=row, column=14, value=metadata.get('breakpoint', ''))
                    ws.cell(row=row, column=15, value=metadata.get('pseudoclass', ''))
                    ws.cell(row=row, column=16, value=page_state_desc)

                    # Add authenticated user info
                    auth_user = metadata.get('authenticated_user', {})
                    if auth_user:
                        ws.cell(row=row, column=17, value=auth_user.get('display_name', ''))
                        ws.cell(row=row, column=18, value=', '.join(auth_user.get('roles', [])))
                    else:
                        ws.cell(row=row, column=17, value='Guest')
                        ws.cell(row=row, column=18, value='no login')

                    for col in range(1, 19):
                        ws.cell(row=row, column=col).fill = styles['discovery']['fill']
                    row += 1

                # Add AI findings if available
                ai_findings = getattr(test_result, 'ai_findings', []) if hasattr(test_result, 'ai_findings') else []
                for f_item in ai_findings:
                    f_dict: dict[str, Any] = self.to_enriched_dict(f_item)

                    ws.cell(row=row, column=1, value='AI Finding')
                    ws.cell(row=row, column=2, value=str(f_dict.get('severity', 'Unknown')).upper())
                    ws.cell(row=row, column=3, value=str(f_dict.get('type', '')))
                    ws.cell(row=row, column=4, value='')  # AI findings don't have touchpoint
                    ws.cell(row=row, column=5, value=str(f_dict.get('description', '')))
                    ws.cell(row=row, column=6, value='')
                    ws.cell(row=row, column=7, value='')
                    ws.cell(row=row, column=8, value=str(f_dict.get('suggested_fix', '')))
                    ws.cell(row=row, column=9, value='')
                    ws.cell(row=row, column=10, value='')
                    ws.cell(row=row, column=11, value='')
                    ws.cell(row=row, column=12, value=page_url)
                    ws.cell(row=row, column=13, value=website_name)
                    ws.cell(row=row, column=14, value='')  # AI findings don't have breakpoint
                    ws.cell(row=row, column=15, value='')  # AI findings don't have pseudoclass
                    ws.cell(row=row, column=16, value=page_state_desc)
                    ws.cell(row=row, column=17, value='')  # AI findings don't have test user
                    ws.cell(row=row, column=18, value='')  # AI findings don't have user roles

                    for col in range(1, 19):
                        ws.cell(row=row, column=col).fill = styles['info']['fill']
                    row += 1

        self._auto_adjust_columns(ws)

    def _xpath_is_within(self, issue_xpath: str, component_xpath: str) -> bool:
        """
        Check if an issue's XPath is within a component's XPath.

        Args:
            issue_xpath: XPath of the issue
            component_xpath: XPath of the component

        Returns:
            True if issue_xpath is within or equal to component_xpath
        """
        if not issue_xpath or not component_xpath:
            return False

        # Normalize xpaths by removing trailing slashes
        issue_xpath = issue_xpath.rstrip('/')
        component_xpath = component_xpath.rstrip('/')

        # Check if issue xpath starts with component xpath
        # e.g., /html/body/nav/a is within /html/body/nav
        return issue_xpath == component_xpath or issue_xpath.startswith(component_xpath + '/')

    def _extract_common_components(self, data: dict[str, Any]) -> dict[str, dict[str, Any]]:
        """
        Extract common components (forms, navs, asides, sections, headers) from discovery issues.

        Args:
            data: Project report data

        Returns:
            Dictionary mapping signature -> component info with xpaths per page
        """
        common_components: dict[str, dict[str, Any]] = {}

        # Iterate through all websites and pages
        for website_data in data.get('websites', []):
            for page_result in website_data.get('pages', []):
                page = page_result.get('page', {})
                page_url: str = self.get_page_url(page)

                test_result = page_result.get('test_result')
                if not test_result:
                    continue

                # Get discovery items
                discovery_items: list[Any] = list(getattr(test_result, 'discovery', []) if hasattr(test_result, 'discovery') else [])

                for d in discovery_items:
                    d_dict: dict[str, Any] = self.to_enriched_dict(d)

                    issue_id: str = str(d_dict.get('id', ''))
                    metadata: dict[str, Any] = d_dict.get('metadata') or {}

                    # Extract signature and xpath for different component types
                    sig: str | None = None
                    component_type: str | None = None
                    label: str | None = None

                    if issue_id in ['DiscoFormOnPage', 'forms_DiscoFormOnPage']:
                        sig = str(metadata.get('formSignature', ''))
                        component_type = 'Form'
                        field_count = metadata.get('fieldCount', 0)
                        label = f"Form ({field_count} fields)"
                    elif issue_id in ['DiscoNavFound', 'landmarks_DiscoNavFound']:
                        sig = str(metadata.get('navSignature', ''))
                        component_type = 'Navigation'
                        label = str(metadata.get('navLabel', 'Navigation'))
                    elif issue_id in ['DiscoAsideFound', 'landmarks_DiscoAsideFound']:
                        sig = str(metadata.get('asideSignature', ''))
                        component_type = 'Aside'
                        label = str(metadata.get('asideLabel', 'Aside'))
                    elif issue_id in ['DiscoSectionFound', 'landmarks_DiscoSectionFound']:
                        sig = str(metadata.get('sectionSignature', ''))
                        component_type = 'Section'
                        label = str(metadata.get('sectionLabel', 'Section'))
                    elif issue_id in ['DiscoHeaderFound', 'landmarks_DiscoHeaderFound']:
                        sig = str(metadata.get('headerSignature', ''))
                        component_type = 'Header'
                        label = str(metadata.get('headerLabel', 'Header'))

                    if sig and sig != 'unknown':
                        if sig not in common_components:
                            common_components[sig] = {
                                'type': component_type,
                                'label': label,
                                'signature': sig,  # Store signature for display
                                'xpaths_by_page': {},  # page_url -> xpath
                                'pages': set()
                            }

                        xpath_val: str = str(d_dict.get('xpath', '') or metadata.get('xpath', ''))
                        common_components[sig]['xpaths_by_page'][page_url] = xpath_val
                        common_components[sig]['pages'].add(page_url)

        return common_components

    def _create_project_deduped_issues_sheet(self, ws: Any, data: dict[str, Any], styles: dict[str, Any]) -> None:
        """Create a deduplicated issues sheet that groups issues by common components"""
        headers = [self._t('type'), self._t('impact'), self._t('rule_id'), self._t('touchpoint'), self._t('what'), self._t('why_important'), self._t('who_affected'),
                   self._t('how_to_remediate'), self._t('wcag_criteria'), self._t('location_xpath'), self._t('element'),
                   self._t('common_component_s'), self._t('pages_with_issue'), self._t('page_count'), self._t('breakpoints'), self._t('pseudoclasses'), self._t('page_states'), self._t('test_users'), self._t('user_roles')]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            self._apply_style(cell, styles['header'])

        row = 2

        # Extract common components from discovery issues
        common_components = self._extract_common_components(data)

        # Track unique issues: (rule_id, xpath_or_component) -> issue data
        unique_issues: dict[tuple[str, str | tuple[str, ...]], dict[str, Any]] = {}

        # Iterate through all websites and their pages
        for website_data in data.get('websites', []):
            for page_result in website_data.get('pages', []):
                page = page_result.get('page', {})
                page_url: str = str(page.get('url', '')) if hasattr(page, 'get') else str(getattr(page, 'url', ''))

                test_result = page_result.get('test_result')
                if not test_result:
                    continue

                # Get page state description from test result if available
                page_state_desc: str = ''
                if test_result:
                    page_state_raw = getattr(test_result, 'page_state', None) if hasattr(test_result, 'page_state') else test_result.get('page_state')
                    if page_state_raw:
                        page_state_desc = str(getattr(page_state_raw, 'description', '') if hasattr(page_state_raw, 'description') else (page_state_raw.get('description', '') if hasattr(page_state_raw, 'get') else ''))

                # Process all issue types
                for issue_type, issue_list_attr in [('violation', 'violations'), ('warning', 'warnings'),
                                                      ('info', 'info'), ('discovery', 'discovery')]:
                    issues: list[Any] = list(getattr(test_result, issue_list_attr, []) if hasattr(test_result, issue_list_attr) else [])

                    for issue in issues:
                        issue_dict: dict[str, Any] = self.to_enriched_dict(issue)

                        rule_id_val: str = str(issue_dict.get('id', ''))
                        issue_xpath: str = str(issue_dict.get('xpath', ''))

                        # Get metadata (breakpoint, pseudoclass)
                        issue_metadata: dict[str, Any] = issue_dict.get('metadata') or {}
                        breakpoint_val: str = str(issue_metadata.get('breakpoint', ''))
                        pseudoclass_val: str = str(issue_metadata.get('pseudoclass', ''))

                        # Find which common components contain this issue
                        containing_components: list[str] = []
                        for _comp_sig, comp_data in common_components.items():
                            # Check if this issue is within this component on this page
                            comp_xpath = comp_data['xpaths_by_page'].get(page_url)
                            if comp_xpath and self._xpath_is_within(issue_xpath, str(comp_xpath)):
                                # Format like Discovery Report: "Type signature"
                                containing_components.append(f"{comp_data['type']} {comp_data['signature']}")

                        # Create deduplication key
                        dedup_key: tuple[str, str | tuple[str, ...]]
                        if containing_components:
                            # Dedupe by rule_id + component(s)
                            dedup_key = (rule_id_val, tuple(sorted(containing_components)))
                        else:
                            # Dedupe by rule_id + exact xpath for non-component issues
                            dedup_key = (rule_id_val, issue_xpath)

                        if dedup_key not in unique_issues:
                            unique_issues[dedup_key] = {
                                'type': issue_type,
                                'data': issue_dict,
                                'component': ', '.join(containing_components) if containing_components else '',
                                'pages': set(),
                                'page_xpaths': {},  # page -> xpath mapping
                                'breakpoints': set(),  # Track all breakpoints where this issue appears
                                'pseudoclasses': set(),  # Track all pseudoclasses
                                'page_states': set(),  # Track all page states
                                'test_users': set(),  # Track all test users who encountered this
                                'user_roles': set()  # Track all unique user roles
                            }

                        unique_issues[dedup_key]['pages'].add(page_url)
                        unique_issues[dedup_key]['page_xpaths'][page_url] = issue_xpath

                        # Add metadata to tracking sets
                        if breakpoint_val:
                            unique_issues[dedup_key]['breakpoints'].add(breakpoint_val)
                        if pseudoclass_val:
                            unique_issues[dedup_key]['pseudoclasses'].add(pseudoclass_val)
                        if page_state_desc:
                            unique_issues[dedup_key]['page_states'].add(page_state_desc)

                        # Track authenticated user info
                        auth_user = issue_metadata.get('authenticated_user', {})
                        if auth_user:
                            user_name = auth_user.get('display_name', '')
                            user_roles_list = auth_user.get('roles', [])

                            if user_name:
                                unique_issues[dedup_key]['test_users'].add(user_name)

                            if user_roles_list:
                                for role in user_roles_list:
                                    unique_issues[dedup_key]['user_roles'].add(role)
                        else:
                            unique_issues[dedup_key]['test_users'].add('Guest')
                            unique_issues[dedup_key]['user_roles'].add('no login')

        # Write deduplicated issues to sheet
        for (rule_id_key, _dedup_value), issue_data in sorted(unique_issues.items(),
                                                          key=lambda x: (x[1]['type'], x[0][0])):
            v_dict = issue_data['data']
            issue_type = issue_data['type']

            # Determine fill color based on type
            if issue_type == 'violation':
                fill_color = styles['violation']['fill']
                type_label = 'Error'
            elif issue_type == 'warning':
                fill_color = styles['warning']['fill']
                type_label = 'Warning'
            elif issue_type == 'info':
                fill_color = styles['info']['fill']
                type_label = 'Info'
            else:  # discovery
                fill_color = styles.get('discovery', {}).get('fill', styles['info']['fill'])
                type_label = 'Discovery'

            ws.cell(row=row, column=1, value=type_label)
            ws.cell(row=row, column=2, value=str(v_dict.get('impact', 'Unknown')).upper())
            ws.cell(row=row, column=3, value=rule_id_key)
            ws.cell(row=row, column=4, value=v_dict.get('touchpoint', v_dict.get('category', '')))
            ws.cell(row=row, column=5, value=v_dict.get('description_full', v_dict.get('what', v_dict.get('description', ''))))
            ws.cell(row=row, column=6, value=v_dict.get('why_it_matters', ''))
            ws.cell(row=row, column=7, value=v_dict.get('who_it_affects', ''))
            ws.cell(row=row, column=8, value=v_dict.get('how_to_fix', v_dict.get('remediation', v_dict.get('suggested_fix', ''))))
            ws.cell(row=row, column=9, value=v_dict.get('wcag_full', ', '.join(v_dict.get('wcag_criteria', [])) if isinstance(v_dict.get('wcag_criteria'), list) else v_dict.get('wcag_criteria', '')))

            # For location, show one representative xpath
            representative_xpath = list(issue_data['page_xpaths'].values())[0] if issue_data['page_xpaths'] else ''
            ws.cell(row=row, column=10, value=representative_xpath)

            ws.cell(row=row, column=11, value=v_dict.get('element', ''))
            ws.cell(row=row, column=12, value=issue_data['component'])

            # List all pages where this issue appears
            pages_list = '\n'.join(sorted(issue_data['pages']))
            ws.cell(row=row, column=13, value=pages_list)
            ws.cell(row=row, column=13).alignment = self.Alignment(wrap_text=True, vertical='top')

            ws.cell(row=row, column=14, value=len(issue_data['pages']))

            # Add metadata columns - show all unique values across all instances of this issue
            breakpoints_list = ', '.join(sorted([str(bp) for bp in issue_data['breakpoints']])) if issue_data['breakpoints'] else ''
            ws.cell(row=row, column=15, value=breakpoints_list)

            pseudoclasses_list = ', '.join(sorted([str(pc) for pc in issue_data['pseudoclasses']])) if issue_data['pseudoclasses'] else ''
            ws.cell(row=row, column=16, value=pseudoclasses_list)

            page_states_list = ', '.join(sorted([str(ps) for ps in issue_data['page_states']])) if issue_data['page_states'] else ''
            ws.cell(row=row, column=17, value=page_states_list)
            ws.cell(row=row, column=17).alignment = self.Alignment(wrap_text=True, vertical='top')

            # Add test user columns
            test_users_list = ', '.join(sorted(issue_data['test_users'])) if issue_data['test_users'] else ''
            ws.cell(row=row, column=18, value=test_users_list)

            user_roles_list = ', '.join(sorted(issue_data['user_roles'])) if issue_data['user_roles'] else ''
            ws.cell(row=row, column=19, value=user_roles_list)

            # Apply fill color
            for col in range(1, 20):
                ws.cell(row=row, column=col).fill = fill_color

            row += 1

        self._auto_adjust_columns(ws)

    def _create_common_components_sheet(self, ws: Any, data: dict[str, Any], styles: dict[str, Any]) -> None:
        """Create a sheet listing all common components identified during deduplication"""
        headers = [self._t('component_type'), self._t('signature'), self._t('label'), self._t('page_count'), self._t('pages_found'), self._t('example_xpath')]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            self._apply_style(cell, styles['header'])

        row = 2

        # Extract common components from discovery issues
        common_components = self._extract_common_components(data)

        # Sort components by type and then by page count (descending)
        sorted_components = sorted(
            common_components.items(),
            key=lambda x: (x[1]['type'], -len(x[1]['pages']))
        )

        for _comp_signature, comp_data in sorted_components:
            # Component type (Navigation, Header, Footer)
            ws.cell(row=row, column=1, value=comp_data['type'])

            # Signature (the unique identifier)
            ws.cell(row=row, column=2, value=comp_data['signature'])

            # Label (display name)
            ws.cell(row=row, column=3, value=comp_data['label'])

            # Page count (moved before Pages Found for better readability)
            ws.cell(row=row, column=4, value=len(comp_data['pages']))

            # Pages found (sorted list)
            pages_list = '\n'.join(sorted(comp_data['pages']))
            ws.cell(row=row, column=5, value=pages_list)
            ws.cell(row=row, column=5).alignment = self.Alignment(wrap_text=True, vertical='top')

            # Example XPath (show one representative xpath)
            example_xpath = list(comp_data['xpaths_by_page'].values())[0] if comp_data['xpaths_by_page'] else ''
            ws.cell(row=row, column=6, value=example_xpath)

            # Apply styling based on component type
            if comp_data['type'] == 'Navigation':
                fill_color = self.PatternFill(start_color="E3F2FD", end_color="E3F2FD", fill_type="solid")
            elif comp_data['type'] == 'Header':
                fill_color = self.PatternFill(start_color="F3E5F5", end_color="F3E5F5", fill_type="solid")
            elif comp_data['type'] == 'Footer':
                fill_color = self.PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
            else:
                fill_color = self.PatternFill(start_color="FFF9C4", end_color="FFF9C4", fill_type="solid")

            for col in range(1, 7):
                ws.cell(row=row, column=col).fill = fill_color

            row += 1

        # Add summary row at the top (after data)
        if sorted_components:
            # Insert a blank row after headers
            ws.insert_rows(2)
            ws.merge_cells('A2:F2')
            summary_cell = ws['A2']
            summary_cell.value = f'Found {len(sorted_components)} common components across {len(data.get("websites", []))} websites'
            summary_cell.font = self.Font(italic=True, color="666666")
            summary_cell.alignment = self.Alignment(horizontal='center')

        self._auto_adjust_columns(ws)

    def _create_website_summary_sheet(self, ws: Any, data: dict[str, Any], styles: dict[str, Any]) -> None:
        """Create website summary sheet"""
        # Title
        ws.merge_cells('A1:D1')
        title_cell = ws['A1']
        title_cell.value = f"{self._t('website_report')} - {data.get('website', {}).get('name', self._t('unknown'))}"
        title_cell.font = self.Font(bold=True, size=14)
        title_cell.alignment = self.Alignment(horizontal="center")

        # Statistics
        row = 3
        stats = data.get('statistics', {})

        ws.cell(row=row, column=1, value=f"{self._t('total_pages')}:").font = self.Font(bold=True)
        ws.cell(row=row, column=2, value=stats.get('total_pages', 0))
        row += 1

        ws.cell(row=row, column=1, value=f"{self._t('total_violations')}:").font = self.Font(bold=True)
        ws.cell(row=row, column=2, value=stats.get('total_violations', 0))
        ws.cell(row=row, column=2).fill = styles['violation']['fill']
        row += 1

        ws.cell(row=row, column=1, value=f"{self._t('total_warnings')}:").font = self.Font(bold=True)
        ws.cell(row=row, column=2, value=stats.get('total_warnings', 0))
        ws.cell(row=row, column=2).fill = styles['warning']['fill']
        row += 1

        ws.cell(row=row, column=1, value=f"{self._t('average_violations_per_page')}:").font = self.Font(bold=True)
        ws.cell(row=row, column=2, value=f"{stats.get('average_violations', 0):.1f}")
        
        self._auto_adjust_columns(ws)
    
    def _create_pages_sheet(self, ws: Any, pages: list[dict[str, Any]], styles: dict[str, Any]) -> None:
        """Create pages sheet"""
        headers = [self._t('page_url'), self._t('violations'), self._t('warnings'), self._t('passes'), self._t('last_tested'), self._t('page_state'), self._t('state_sequence'), self._t('session_id')]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            self._apply_style(cell, styles['header'])

        row = 2
        for pr in pages:
            page = pr.get('page', {})
            test = pr.get('test_result', {})

            ws.cell(row=row, column=1, value=page.get('url', ''))

            def _test_val(key: str, default: Any = '') -> Any:
                """Get value from test result dict or object."""
                return test.get(key, default) if hasattr(test, 'get') else getattr(test, key, default)

            violation_count: int = int(_test_val('violation_count', 0))
            violations_cell = ws.cell(row=row, column=2, value=violation_count)
            if violation_count > 0:
                violations_cell.fill = styles['violation']['fill']

            warning_count: int = int(_test_val('warning_count', 0))
            warnings_cell = ws.cell(row=row, column=3, value=warning_count)
            if warning_count > 0:
                warnings_cell.fill = styles['warning']['fill']

            pass_count: int = int(_test_val('pass_count', 0))
            ws.cell(row=row, column=4, value=pass_count)
            test_date: str = str(_test_val('test_date', ''))
            ws.cell(row=row, column=5, value=test_date)

            # Add multi-state information
            page_state_raw = _test_val('page_state', None)
            if page_state_raw:
                state_desc_val: str = str(getattr(page_state_raw, 'description', '') if hasattr(page_state_raw, 'description') else (page_state_raw.get('description', '') if hasattr(page_state_raw, 'get') else ''))
                ws.cell(row=row, column=6, value=state_desc_val)
            else:
                ws.cell(row=row, column=6, value='')

            state_seq_val: str = str(_test_val('state_sequence', ''))
            ws.cell(row=row, column=7, value=state_seq_val if state_seq_val != '0' else '')

            session_id_val: str = str(_test_val('session_id', ''))
            ws.cell(row=row, column=8, value=session_id_val or '')

            row += 1

        self._auto_adjust_columns(ws)
    
    def _create_violation_types_sheet(self, ws: Any, violation_types: dict[str, Any], styles: dict[str, Any]) -> None:
        """Create violation types sheet"""
        headers = [self._t('rule_id'), self._t('count'), self._t('description'), self._t('pages_affected')]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            self._apply_style(cell, styles['header'])
        
        row = 2
        sorted_types = sorted(violation_types.items(), key=lambda x: x[1]['count'], reverse=True)
        
        for rule_id, info in sorted_types:
            ws.cell(row=row, column=1, value=rule_id)
            ws.cell(row=row, column=2, value=info['count'])
            ws.cell(row=row, column=3, value=info.get('description', ''))
            ws.cell(row=row, column=4, value=len(set(info.get('pages', []))))
            row += 1
        
        self._auto_adjust_columns(ws)
    
    def _create_project_summary_sheet(self, ws: Any, data: dict[str, Any], styles: dict[str, Any]) -> None:
        """Create project summary sheet"""
        # Title
        ws.merge_cells('A1:D1')
        title_cell = ws['A1']
        title_cell.value = f"{self._t('project_report')} - {data.get('project', {}).get('name', self._t('unknown'))}"
        title_cell.font = self.Font(bold=True, size=14)
        title_cell.alignment = self.Alignment(horizontal="center")

        # Statistics
        row = 3
        stats = data.get('statistics', {})

        ws.cell(row=row, column=1, value=f"{self._t('total_websites')}:").font = self.Font(bold=True)
        ws.cell(row=row, column=2, value=stats.get('total_websites', 0))
        row += 1

        ws.cell(row=row, column=1, value=f"{self._t('total_pages')}:").font = self.Font(bold=True)
        ws.cell(row=row, column=2, value=stats.get('total_pages', 0))
        row += 1

        ws.cell(row=row, column=1, value=f"{self._t('total_violations')}:").font = self.Font(bold=True)
        ws.cell(row=row, column=2, value=stats.get('total_violations', 0))
        ws.cell(row=row, column=2).fill = styles['violation']['fill']
        row += 1

        ws.cell(row=row, column=1, value=f"{self._t('average_violations_per_page')}:").font = self.Font(bold=True)
        ws.cell(row=row, column=2, value=f"{stats.get('average_violations_per_page', 0):.1f}")
        
        self._auto_adjust_columns(ws)
    
    def _create_websites_sheet(self, ws: Any, websites: list[dict[str, Any]], styles: dict[str, Any]) -> None:
        """Create websites sheet"""
        headers = [self._t('website_name'), self._t('url'), self._t('pages'), self._t('total_violations'), self._t('total_warnings')]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            self._apply_style(cell, styles['header'])
        
        row = 2
        for wd in websites:
            website = wd.get('website', {})
            pages = wd.get('pages', [])
            
            # Handle test_result as object or dict
            total_violations: int = 0
            total_warnings: int = 0
            for p in pages:
                test_result = p.get('test_result')
                if test_result:
                    if hasattr(test_result, 'violation_count'):
                        # It's an object
                        total_violations += int(test_result.violation_count)
                        total_warnings += int(test_result.warning_count)
                    elif hasattr(test_result, 'get'):
                        # It's a dictionary
                        total_violations += int(test_result.get('violation_count', 0))
                        total_warnings += int(test_result.get('warning_count', 0))
            
            # Handle website as object or dict
            if hasattr(website, 'name'):
                # It's an object
                ws.cell(row=row, column=1, value=website.name or '')
                ws.cell(row=row, column=2, value=getattr(website, 'url', None) or getattr(website, 'base_url', ''))
            elif hasattr(website, 'get'):
                # It's a dictionary
                ws.cell(row=row, column=1, value=str(website.get('name', '')))
                ws.cell(row=row, column=2, value=str(website.get('url', website.get('base_url', ''))))
            ws.cell(row=row, column=3, value=len(pages))
            
            violations_cell = ws.cell(row=row, column=4, value=total_violations)
            if total_violations > 0:
                violations_cell.fill = styles['violation']['fill']
            
            warnings_cell = ws.cell(row=row, column=5, value=total_warnings)
            if total_warnings > 0:
                warnings_cell.fill = styles['warning']['fill']
            
            row += 1
        
        self._auto_adjust_columns(ws)
    
    def _create_all_projects_summary_sheet(self, ws: Any, data: dict[str, Any], styles: dict[str, Any]) -> None:
        """Create all projects summary sheet"""
        # Title
        ws.merge_cells('A1:F1')
        title_cell = ws['A1']
        title_cell.value = self._t('all_projects_accessibility_report')
        title_cell.font = self.Font(bold=True, size=14)
        title_cell.alignment = self.Alignment(horizontal="center")

        # Overall Statistics
        row = 3
        summary = data.get('summary', {})

        ws.cell(row=row, column=1, value=self._t('overall_statistics')).font = self.Font(bold=True, size=12)
        row += 1

        ws.cell(row=row, column=1, value=f"{self._t('total_projects')}:").font = self.Font(bold=True)
        ws.cell(row=row, column=2, value=summary.get('total_projects', 0))
        row += 1

        ws.cell(row=row, column=1, value=f"{self._t('total_websites')}:").font = self.Font(bold=True)
        ws.cell(row=row, column=2, value=summary.get('total_websites', 0))
        row += 1

        ws.cell(row=row, column=1, value=f"{self._t('total_pages')}:").font = self.Font(bold=True)
        ws.cell(row=row, column=2, value=summary.get('total_pages', 0))
        row += 1

        ws.cell(row=row, column=1, value=f"{self._t('pages_tested')}:").font = self.Font(bold=True)
        ws.cell(row=row, column=2, value=summary.get('total_tested', 0))
        row += 1

        ws.cell(row=row, column=1, value=f"{self._t('total_violations')}:").font = self.Font(bold=True)
        ws.cell(row=row, column=2, value=summary.get('total_violations', 0))
        ws.cell(row=row, column=2).fill = styles['violation']['fill']
        row += 1

        ws.cell(row=row, column=1, value=f"{self._t('total_warnings')}:").font = self.Font(bold=True)
        ws.cell(row=row, column=2, value=summary.get('total_warnings', 0))
        ws.cell(row=row, column=2).fill = styles['warning']['fill']
        
        self._auto_adjust_columns(ws)
    
    def _create_projects_breakdown_sheet(self, ws: Any, projects: list[dict[str, Any]], styles: dict[str, Any]) -> None:
        """Create projects breakdown sheet"""
        headers = [self._t('project_name'), self._t('description'), self._t('websites'), self._t('total_pages'), self._t('tested_pages'), self._t('coverage_pct'), self._t('violations'), self._t('warnings')]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            self._apply_style(cell, styles['header'])
        
        row = 2
        for project_data in projects:
            project = project_data.get('project', {})
            stats = project_data.get('stats', {})
            
            ws.cell(row=row, column=1, value=project.get('name', ''))
            ws.cell(row=row, column=2, value=project.get('description', ''))
            ws.cell(row=row, column=3, value=stats.get('website_count', 0))
            ws.cell(row=row, column=4, value=stats.get('total_pages', 0))
            ws.cell(row=row, column=5, value=stats.get('tested_pages', 0))
            ws.cell(row=row, column=6, value=f"{stats.get('test_coverage', 0):.1f}%")
            
            violations = sum(w.get('violations', 0) for w in project_data.get('websites', []))
            violations_cell = ws.cell(row=row, column=7, value=violations)
            if violations > 0:
                violations_cell.fill = styles['violation']['fill']
            
            warnings = sum(w.get('warnings', 0) for w in project_data.get('websites', []))
            warnings_cell = ws.cell(row=row, column=8, value=warnings)
            if warnings > 0:
                warnings_cell.fill = styles['warning']['fill']
            
            row += 1
        
        self._auto_adjust_columns(ws)
    
    def _apply_style(self, cell: Any, style: dict[str, Any]) -> None:
        """Apply style dictionary to a cell"""
        for attr, value in style.items():
            setattr(cell, attr, value)
    
    def _create_page_states_sheet(self, ws: Any, data: dict[str, Any], styles: dict[str, Any]) -> None:
        """Create sheet showing multi-state test results summary"""
        # Title
        ws.merge_cells('A1:G1')
        title_cell = ws['A1']
        title_cell.value = self._t('page_state_information')
        title_cell.font = self.Font(bold=True, size=14)
        title_cell.alignment = self.Alignment(horizontal="center")

        # Headers
        row = 3
        headers = [self._t('state_sequence'), self._t('state_description'), self._t('errors'), self._t('warnings'), self._t('info'), self._t('discovery'), self._t('test_date')]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=header)
            self._apply_style(cell, styles['header'])

        row += 1

        # Get test result
        test_result = data.get('test_result', {})
        session_id = test_result.get('session_id')

        # Show the current state
        state_sequence = test_result.get('state_sequence', 0)
        page_state = test_result.get('page_state', {})

        page_state_dict: dict[str, Any] = page_state if hasattr(page_state, 'get') else {}
        state_desc: str = str(page_state_dict.get('description', f'State {state_sequence}'))
        scripts_executed: list[Any] = list(page_state_dict.get('scripts_executed', []) or [])
        elements_clicked: list[Any] = list(page_state_dict.get('elements_clicked', []) or [])

        ws.cell(row=row, column=1, value=state_sequence)
        ws.cell(row=row, column=2, value=state_desc)
        ws.cell(row=row, column=3, value=len(data.get('violations', [])))
        ws.cell(row=row, column=4, value=len(data.get('warnings', [])))
        ws.cell(row=row, column=5, value=len(data.get('info', [])))
        ws.cell(row=row, column=6, value=len(data.get('discovery', [])))
        ws.cell(row=row, column=7, value=str(test_result.get('test_date', '')))

        # Add state details section
        row += 2
        ws.cell(row=row, column=1, value=f"{self._t('state_details')}:").font = self.Font(bold=True, size=12)
        row += 1

        if session_id:
            ws.cell(row=row, column=1, value=f"{self._t('session_id')}:").font = self.Font(bold=True)
            ws.cell(row=row, column=2, value=session_id)
            row += 1

        if scripts_executed:
            ws.cell(row=row, column=1, value=f"{self._t('scripts_executed')}:").font = self.Font(bold=True)
            ws.cell(row=row, column=2, value=', '.join(str(s) for s in scripts_executed))
            row += 1

        if elements_clicked:
            ws.cell(row=row, column=1, value=f"{self._t('elements_clicked')}:").font = self.Font(bold=True)
            if len(elements_clicked) > 0:
                click_desc = ', '.join(str(el_item.get('description', el_item.get('selector', str(el_item)))) if hasattr(el_item, 'get') else str(el_item) for el_item in elements_clicked)
                ws.cell(row=row, column=2, value=click_desc)
            row += 1

        # Add note about multi-state testing
        row += 1
        ws.cell(row=row, column=1, value=f"{self._t('note')}:").font = self.Font(bold=True)
        row += 1
        ws.cell(row=row, column=1, value=f"{self._t('multistate_testing')}:")
        ws.cell(row=row, column=2, value=self._t('multistate_note'))
        row += 1
        ws.cell(row=row, column=1, value=f"{self._t('breakpoint_testing')}:")
        ws.cell(row=row, column=2, value=self._t('breakpoint_note'))
        row += 1
        ws.cell(row=row, column=1, value=f"{self._t('context_information')}:")
        ws.cell(row=row, column=2, value=self._t('context_note'))

        self._auto_adjust_columns(ws)

    def _auto_adjust_columns(self, ws: Any) -> None:
        """Auto-adjust column widths"""
        for column in ws.columns:
            max_length = 0
            column_letter = self.get_column_letter(column[0].column)

            for cell in column:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except Exception as exc:
                    logger.debug("Could not measure column width for cell: %s", exc)
            
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

    # --- Streaming interface ---

    def _detail_headers(self) -> list[str]:
        """Get translated detail headers for streaming sheets"""
        return [self._t('page_url'), self._t('page_title'), self._t('code'), self._t('description'),
                self._t('touchpoint'), self._t('impact'), self._t('xpath'), self._t('html'), self._t('wcag_criteria')]

    @override
    def begin(self, output_file: str, summary: dict[str, Any]) -> None:
        """Create a Workbook with Summary, Violations, and Warnings sheets."""
        self._output_file = output_file
        self._wb = self.Workbook()
        styles = self._get_styles()

        # --- Summary sheet ---
        ws_sum = self._wb.active
        assert ws_sum is not None
        ws_sum.title = self._t('summary')
        ws_sum.merge_cells('A1:D1')
        hdr_cell = ws_sum.cell(row=1, column=1, value=self._t('accessibility_report_summary'))
        for attr in ('font', 'fill', 'alignment'):
            setattr(hdr_cell, attr, styles['header'][attr])

        row = 3
        for key, val in (summary or {}).items():
            ws_sum.cell(row=row, column=1, value=str(key)).font = self.Font(bold=True)
            ws_sum.cell(row=row, column=2, value=str(val))
            row += 1

        # --- Violations detail sheet ---
        self._ws_violations = self._wb.create_sheet(self._t('violations'))
        for col_idx, hdr in enumerate(self._detail_headers(), 1):
            cell = self._ws_violations.cell(row=1, column=col_idx, value=hdr)
            for attr in ('font', 'fill', 'alignment'):
                setattr(cell, attr, styles['header'][attr])

        # --- Warnings detail sheet ---
        self._ws_warnings = self._wb.create_sheet(self._t('warnings'))
        for col_idx, hdr in enumerate(self._detail_headers(), 1):
            cell = self._ws_warnings.cell(row=1, column=col_idx, value=hdr)
            for attr in ('font', 'fill', 'alignment'):
                setattr(cell, attr, styles['header'][attr])

    @override
    def append_page(self, output_file: str, page_data: dict[str, Any]) -> None:
        """Append rows for one page to the Violations and Warnings sheets."""
        if self._wb is None:
            raise RuntimeError("ExcelFormatter.append_page() called before begin()")

        page = page_data.get('page', {})
        test_result = page_data.get('test_result')
        if test_result is None:
            return

        page_url: str = self.get_page_url(page)
        page_title: str = self.get_page_title(page)

        violations: list[Any] = self.get_issue_list(test_result, 'violations')
        warnings_list: list[Any] = self.get_issue_list(test_result, 'warnings')

        for v in violations:
            self._append_issue_row(self._ws_violations, page_url, page_title, v)
        for w in warnings_list:
            self._append_issue_row(self._ws_warnings, page_url, page_title, w)

    @override
    def finalize(self, output_file: str, summary: dict[str, Any]) -> None:
        """Auto-size columns and save the workbook to *output_file*."""
        if self._wb is None:
            raise RuntimeError("ExcelFormatter.finalize() called before begin()")

        for ws in self._wb.worksheets:
            self._auto_adjust_columns(ws)

        self._wb.save(output_file)

    @override
    def cleanup(self) -> None:
        """Close the workbook if still open."""
        if hasattr(self, '_wb') and self._wb:
            try:
                self._wb.close()
            except Exception:
                pass

    # --- streaming helpers ---

    def _append_issue_row(self, ws: Any, page_url: str, page_title: str, issue: Any) -> None:
        """Append a single data row to a worksheet."""
        issue_dict: dict[str, Any] = self.to_enriched_dict(issue)

        impact_raw = issue_dict.get('impact', '')
        if hasattr(impact_raw, 'value'):
            impact_raw = impact_raw.value

        wcag_str: str = self._format_wcag(issue_dict)

        ws.append([
            page_url,
            page_title,
            str(issue_dict.get('id', '')),
            self._best_description(issue_dict),
            str(issue_dict.get('touchpoint', '')),
            self._translate_impact(str(impact_raw)),
            issue_dict.get('xpath', ''),
            issue_dict.get('html', ''),
            wcag_str,
        ])


class PDFFormatter(BaseFormatter):
    """PDF report formatter (uses HTML + conversion)"""

    def __init__(self, config: dict[str, Any], language: str = 'en') -> None:
        super().__init__(config, language)
        self.extension = 'pdf'
        self.html_formatter = HTMLFormatter(config, language)
        # Streaming state
        self._pdf_output_file: str = ''
        self._internal_html: HTMLFormatter | None = None
        self._temp_html_path: str = ''
        from weasyprint import HTML, CSS
        self.HTML = HTML
        self.CSS = CSS

    def _wrap_render_error(self, err: BaseException) -> RuntimeError:
        """Wrap a weasyprint runtime error with an actionable message."""
        message = f"PDF report generation failed: {err}"
        if is_native_lib_error(err):
            message = f"{message}\n\n{platform_install_hint()}"
        return RuntimeError(message)
    
    @override
    def format_page_report(self, data: dict[str, Any]) -> bytes:
        """Generate PDF for page report"""
        html_content = self.html_formatter.format_page_report(data)
        return self._convert_to_pdf(html_content)

    @override
    def format_website_report(self, data: dict[str, Any]) -> bytes:
        """Generate PDF for website report"""
        html_content = self.html_formatter.format_website_report(data)
        return self._convert_to_pdf(html_content)

    @override
    def format_project_report(self, data: dict[str, Any]) -> bytes:
        """Generate PDF for project report"""
        html_content = self.html_formatter.format_project_report(data)
        return self._convert_to_pdf(html_content)
    
    def format_all_projects_report(self, data: dict[str, Any]) -> bytes:
        """Generate PDF for all projects report"""
        warnings.warn(
            "format_all_projects_report() is deprecated, use begin/append_page/finalize streaming interface",
            DeprecationWarning,
            stacklevel=2
        )
        html_content = self.html_formatter.format_all_projects_report(data)
        return self._convert_to_pdf(html_content)
    
    @override
    def format_summary_report(self, data: dict[str, Any]) -> bytes:
        """Generate PDF for summary report"""
        html_content = self.html_formatter.format_summary_report(data)
        return self._convert_to_pdf(html_content)
    
    def _convert_to_pdf(self, html_content: str) -> bytes:
        """Convert HTML content to PDF bytes"""
        pdf_css = self.CSS(string='''
            @page {
                size: A4;
                margin: 1cm;
            }
            body {
                font-size: 10pt;
            }
            .container {
                max-width: 100%;
                box-shadow: none;
            }
            table {
                page-break-inside: avoid;
            }
            .violation, .warning, .pass, .ai-finding {
                page-break-inside: avoid;
            }
        ''')

        try:
            pdf_document = self.HTML(string=html_content).render(stylesheets=[pdf_css])
            write_fn = getattr(pdf_document, 'write_pdf')
            pdf_bytes: bytes = write_fn()
            return pdf_bytes
        except Exception as e:
            logger.error(f"Failed to generate PDF with weasyprint: {e}", exc_info=True)
            raise self._wrap_render_error(e) from e
    
    def save_pdf(self, html_content: str, filepath: Path) -> None:
        """
        Save HTML as PDF

        This method is deprecated - use the format methods that return bytes instead
        """
        pdf_bytes = self._convert_to_pdf(html_content)
        with open(filepath, 'wb') as f:
            f.write(pdf_bytes)

    # --- Streaming interface ---

    @override
    def begin(self, output_file: str, summary: dict[str, Any]) -> None:
        """Create an internal HTMLFormatter and a temp HTML file, then delegate."""
        self._pdf_output_file = output_file
        self._internal_html = HTMLFormatter(self.config, self.language)
        # PDF keeps the per-page detail section and never triggers the AI
        # call — the AI executive analysis is a feature of the standalone
        # HTML report only.
        self._internal_html.include_page_detail = True
        self._internal_html.include_ai_summary = False
        # Temp HTML file that the internal formatter writes to
        fd, self._temp_html_path = tempfile.mkstemp(suffix='.html')
        os.close(fd)
        self._internal_html.begin(self._temp_html_path, summary)

    @override
    def append_page(self, output_file: str, page_data: dict[str, Any]) -> None:
        """Delegate to internal HTMLFormatter."""
        if hasattr(self, '_internal_html') and self._internal_html:
            self._internal_html.append_page(self._temp_html_path, page_data)

    @override
    def finalize(self, output_file: str, summary: dict[str, Any]) -> None:
        """Finalize the internal HTML, then convert to PDF via weasyprint."""
        if not hasattr(self, '_internal_html') or not self._internal_html:
            raise RuntimeError("PDFFormatter.finalize() called before begin()")

        self._internal_html.finalize(self._temp_html_path, summary)

        try:
            self.HTML(filename=self._temp_html_path).write_pdf(output_file)
        except Exception as e:
            logger.error(f"PDF streaming conversion failed: {e}", exc_info=True)
            raise self._wrap_render_error(e) from e

    @override
    def cleanup(self) -> None:
        """Clean up internal HTMLFormatter temps and own temp HTML file."""
        if hasattr(self, '_internal_html') and self._internal_html:
            self._internal_html.cleanup()
        if hasattr(self, '_temp_html_path') and self._temp_html_path:
            if os.path.exists(self._temp_html_path):
                os.unlink(self._temp_html_path)