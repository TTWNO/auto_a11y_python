"""Tests for streaming interface on report formatters."""
from __future__ import annotations

from typing import Any

import json
import csv
import os
import tempfile
import shutil
import pytest

from auto_a11y.reporting.formatters import (
    BaseFormatter, CSVFormatter, JSONFormatter, HTMLFormatter, ExcelFormatter,
    PDFFormatter,
)


# ---------------------------------------------------------------------------
# Task 2 -- BaseFormatter streaming interface
# ---------------------------------------------------------------------------

class TestBaseFormatterStreaming:
    """Verify the abstract streaming methods on BaseFormatter."""

    formatter = BaseFormatter(config={})

    def setup_method(self) -> None:
        self.formatter = BaseFormatter(config={})

    def test_begin_raises(self) -> None:
        with pytest.raises(NotImplementedError):
            self.formatter.begin("out.txt", {})

    def test_append_page_raises(self) -> None:
        with pytest.raises(NotImplementedError):
            self.formatter.append_page("out.txt", {})

    def test_finalize_raises(self) -> None:
        with pytest.raises(NotImplementedError):
            self.formatter.finalize("out.txt", {})

    def test_cleanup_is_noop(self) -> None:
        # Should not raise (implicitly returns None)
        self.formatter.cleanup()


# ---------------------------------------------------------------------------
# Helpers -- lightweight stand-ins for model objects
# ---------------------------------------------------------------------------

class _FakePage:
    def __init__(self, url: str = 'http://example.com', title: str = 'Example') -> None:
        self.url = url
        self.title = title


class _FakeIssue:
    def __init__(self, **kwargs: Any) -> None:
        self.id: str = kwargs.get('id', 'ErrTest')
        self.description: str = kwargs.get('description', 'desc')
        self.touchpoint: str = kwargs.get('touchpoint', 'Forms')
        self.impact: str = kwargs.get('impact', 'critical')
        self.xpath: str = kwargs.get('xpath', '/html/body')
        self.html: str = kwargs.get('html', '<input>')
        self.wcag_criteria: list[str] = kwargs.get('wcag_criteria', ['1.1.1'])
        self.metadata: dict[str, Any] = kwargs.get('metadata', {})


class _FakeTestResult:
    def __init__(self, violations: list[Any] | None = None,
                 warnings: list[Any] | None = None) -> None:
        self.violations = violations or []
        self.warnings = warnings or []


# ---------------------------------------------------------------------------
# Task 3 -- CSVFormatter streaming
# ---------------------------------------------------------------------------

class TestCSVFormatterStreaming:
    """CSVFormatter.begin / append_page / finalize / cleanup."""

    tmpdir = ""
    outfile = ""
    formatter = CSVFormatter(config={})

    def setup_method(self) -> None:
        self.tmpdir = tempfile.mkdtemp()
        self.outfile = os.path.join(self.tmpdir, 'report.csv')
        self.formatter = CSVFormatter(config={})

    def teardown_method(self) -> None:
        self.formatter.cleanup()
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def _read_rows(self) -> list[list[str]]:
        with open(self.outfile, newline='', encoding='utf-8') as f:
            return list(csv.reader(f))

    def test_header_row_written(self) -> None:
        self.formatter.begin(self.outfile, {})
        self.formatter.finalize(self.outfile, {})
        rows = self._read_rows()
        assert len(rows) == 1
        assert rows[0] == [
            'URL', 'Page Title', 'Type', 'Code', 'Description',
            'Touchpoint', 'Impact', 'XPath', 'HTML', 'WCAG Criteria',
        ]

    def test_violation_rows_written(self) -> None:
        page_data: dict[str, Any] = {
            'page': _FakePage(url='http://a.com', title='A'),
            'test_result': _FakeTestResult(
                violations=[_FakeIssue(id='ErrNoAlt', description='Missing alt')],
                warnings=[_FakeIssue(id='WarnContrast', description='Low contrast')],
            ),
        }
        self.formatter.begin(self.outfile, {})
        self.formatter.append_page(self.outfile, page_data)
        self.formatter.finalize(self.outfile, {})

        rows = self._read_rows()
        # header + 1 violation + 1 warning = 3
        assert len(rows) == 3
        assert rows[1][0] == 'http://a.com'  # URL
        assert rows[1][2] == 'Violation'
        assert rows[1][3] == 'ErrNoAlt'
        assert rows[2][2] == 'Warning'
        assert rows[2][3] == 'WarnContrast'

    def test_multiple_pages_accumulate(self) -> None:
        self.formatter.begin(self.outfile, {})
        for i in range(3):
            page_data: dict[str, Any] = {
                'page': _FakePage(url=f'http://p{i}.com', title=f'P{i}'),
                'test_result': _FakeTestResult(
                    violations=[_FakeIssue()],
                ),
            }
            self.formatter.append_page(self.outfile, page_data)
        self.formatter.finalize(self.outfile, {})

        rows = self._read_rows()
        assert len(rows) == 4  # header + 3 violation rows

    def test_dict_page_data(self) -> None:
        """page_data may use plain dicts instead of model objects."""
        page_data: dict[str, Any] = {
            'page': {'url': 'http://dict.com', 'title': 'Dict Page'},
            'test_result': {
                'violations': [{'id': 'ErrX', 'description': 'd', 'touchpoint': 'T',
                                'impact': 'serious', 'xpath': '/x', 'html': '<b>',
                                'wcag_criteria': ['4.1.2']}],
                'warnings': [],
            },
        }
        self.formatter.begin(self.outfile, {})
        self.formatter.append_page(self.outfile, page_data)
        self.formatter.finalize(self.outfile, {})

        rows = self._read_rows()
        assert len(rows) == 2
        assert rows[1][0] == 'http://dict.com'
        assert rows[1][3] == 'ErrX'

    def test_cleanup_closes_file(self) -> None:
        self.formatter.begin(self.outfile, {})
        self.formatter.cleanup()
        file_obj = getattr(self.formatter, '_file')
        assert file_obj is not None and file_obj.closed


# ---------------------------------------------------------------------------
# Task 4 -- JSONFormatter streaming
# ---------------------------------------------------------------------------

class TestJSONFormatterStreaming:
    """JSONFormatter.begin / append_page / finalize / cleanup."""

    tmpdir = ""
    outfile = ""
    formatter = JSONFormatter(config={})

    def setup_method(self) -> None:
        self.tmpdir = tempfile.mkdtemp()
        self.outfile = os.path.join(self.tmpdir, 'report.json')
        self.formatter = JSONFormatter(config={})

    def teardown_method(self) -> None:
        self.formatter.cleanup()
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def _load_json(self) -> Any:
        with open(self.outfile, encoding='utf-8') as f:
            return json.load(f)

    def test_valid_json_with_summary_and_pages(self) -> None:
        summary: dict[str, Any] = {'total_pages': 1, 'total_violations': 2}
        page_data: dict[str, Any] = {
            'page': _FakePage(url='http://a.com', title='A'),
            'test_result': _FakeTestResult(
                violations=[_FakeIssue(id='ErrNoAlt')],
            ),
        }
        self.formatter.begin(self.outfile, summary)
        self.formatter.append_page(self.outfile, page_data)
        self.formatter.finalize(self.outfile, summary)

        data = self._load_json()
        assert data['summary']['total_pages'] == 1
        assert len(data['pages']) == 1
        assert data['pages'][0]['page']['url'] == 'http://a.com'

    def test_multiple_pages_valid_json(self) -> None:
        summary: dict[str, Any] = {'n': 3}
        self.formatter.begin(self.outfile, summary)
        for i in range(3):
            page_data: dict[str, Any] = {
                'page': _FakePage(url=f'http://p{i}.com', title=f'P{i}'),
                'test_result': _FakeTestResult(violations=[_FakeIssue()]),
            }
            self.formatter.append_page(self.outfile, page_data)
        self.formatter.finalize(self.outfile, summary)

        data = self._load_json()
        assert len(data['pages']) == 3
        urls = [p['page']['url'] for p in data['pages']]
        assert urls == ['http://p0.com', 'http://p1.com', 'http://p2.com']

    def test_zero_pages_valid_json(self) -> None:
        summary: dict[str, Any] = {'empty': True}
        self.formatter.begin(self.outfile, summary)
        self.formatter.finalize(self.outfile, summary)

        data = self._load_json()
        assert data['summary'] == {'empty': True}
        assert data['pages'] == []

    def test_cleanup_closes_file(self) -> None:
        self.formatter.begin(self.outfile, {})
        self.formatter.cleanup()
        file_obj = getattr(self.formatter, '_file')
        assert file_obj is not None and file_obj.closed

    def test_dict_page_data(self) -> None:
        """page_data with plain dicts instead of model objects."""
        page_data: dict[str, Any] = {
            'page': {'url': 'http://dict.com', 'title': 'Dict'},
            'test_result': {
                'violations': [{'id': 'ErrX', 'description': 'd'}],
                'warnings': [],
            },
        }
        self.formatter.begin(self.outfile, {})
        self.formatter.append_page(self.outfile, page_data)
        self.formatter.finalize(self.outfile, {})

        data = self._load_json()
        assert len(data['pages']) == 1
        assert data['pages'][0]['page']['url'] == 'http://dict.com'


# ---------------------------------------------------------------------------
# Task 5 -- HTMLFormatter streaming
# ---------------------------------------------------------------------------

class TestHTMLFormatterStreaming:
    """HTMLFormatter.begin / append_page / finalize / cleanup."""

    tmpdir = ""
    outfile = ""
    formatter = HTMLFormatter(config={})

    def setup_method(self) -> None:
        self.tmpdir = tempfile.mkdtemp()
        self.outfile = os.path.join(self.tmpdir, 'report.html')
        self.formatter = HTMLFormatter(config={})

    def teardown_method(self) -> None:
        self.formatter.cleanup()
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_produces_valid_html_document(self) -> None:
        """begin/append/finalize produces HTML with <html> and </html>."""
        summary: dict[str, Any] = {'total_pages': 1, 'total_violations': 1}
        page_data: dict[str, Any] = {
            'page': _FakePage(url='http://a.com', title='A'),
            'test_result': _FakeTestResult(
                violations=[_FakeIssue(id='ErrNoAlt', description='Missing alt')],
            ),
        }
        self.formatter.begin(self.outfile, summary)
        self.formatter.append_page(self.outfile, page_data)
        self.formatter.finalize(self.outfile, summary)

        with open(self.outfile, encoding='utf-8') as f:
            content = f.read()
        assert '<html' in content
        assert '</html>' in content

    def test_cleanup_removes_temp_file(self) -> None:
        """After begin(), temp file exists; after cleanup(), it doesn't."""
        self.formatter.begin(self.outfile, {})
        body_tempfile = getattr(self.formatter, '_body_tempfile')
        assert body_tempfile is not None
        temp_path = body_tempfile.name
        assert os.path.exists(temp_path)

        self.formatter.cleanup()
        assert not os.path.exists(temp_path)

    def test_body_content_produces_nonempty_file(self) -> None:
        """Output file exists and is non-empty after full cycle."""
        summary: dict[str, Any] = {'pages': 2}
        self.formatter.begin(self.outfile, summary)
        for i in range(2):
            page_data: dict[str, Any] = {
                'page': _FakePage(url=f'http://p{i}.com', title=f'P{i}'),
                'test_result': _FakeTestResult(
                    violations=[_FakeIssue()],
                    warnings=[_FakeIssue(id='WarnC')],
                ),
            }
            self.formatter.append_page(self.outfile, page_data)
        self.formatter.finalize(self.outfile, summary)

        assert os.path.exists(self.outfile)
        assert os.path.getsize(self.outfile) > 0


# ---------------------------------------------------------------------------
# Task 6 -- ExcelFormatter streaming
# ---------------------------------------------------------------------------

class TestExcelFormatterStreaming:
    """ExcelFormatter.begin / append_page / finalize / cleanup."""

    tmpdir = ""
    outfile = ""
    formatter = ExcelFormatter(config={})

    def setup_method(self) -> None:
        self.tmpdir = tempfile.mkdtemp()
        self.outfile = os.path.join(self.tmpdir, 'report.xlsx')
        self.formatter = ExcelFormatter(config={})

    def teardown_method(self) -> None:
        self.formatter.cleanup()
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    @pytest.mark.skipif(
        not ExcelFormatter(config={}).has_openpyxl,
        reason='openpyxl not installed',
    )
    def test_produces_valid_xlsx(self) -> None:
        """Produces valid xlsx with at least 2 sheets."""
        from openpyxl import load_workbook

        summary: dict[str, Any] = {'total_pages': 1, 'total_violations': 2}
        page_data: dict[str, Any] = {
            'page': _FakePage(url='http://a.com', title='A'),
            'test_result': _FakeTestResult(
                violations=[_FakeIssue(id='ErrNoAlt')],
                warnings=[_FakeIssue(id='WarnC')],
            ),
        }
        self.formatter.begin(self.outfile, summary)
        self.formatter.append_page(self.outfile, page_data)
        self.formatter.finalize(self.outfile, summary)

        wb = load_workbook(self.outfile)
        assert len(wb.sheetnames) >= 2
        wb.close()

    @pytest.mark.skipif(
        not ExcelFormatter(config={}).has_openpyxl,
        reason='openpyxl not installed',
    )
    def test_multiple_pages_add_rows(self) -> None:
        """Violations sheet has header + rows from multiple pages."""
        from openpyxl import load_workbook

        self.formatter.begin(self.outfile, {})
        for i in range(3):
            page_data: dict[str, Any] = {
                'page': _FakePage(url=f'http://p{i}.com', title=f'P{i}'),
                'test_result': _FakeTestResult(
                    violations=[_FakeIssue(id=f'Err{i}')],
                ),
            }
            self.formatter.append_page(self.outfile, page_data)
        self.formatter.finalize(self.outfile, {})

        wb = load_workbook(self.outfile)
        ws = wb['Violations']
        # 1 header row + 3 data rows = 4
        assert ws.max_row == 4
        wb.close()


# ---------------------------------------------------------------------------
# Task 7 -- PDFFormatter streaming
# ---------------------------------------------------------------------------

def _has_weasyprint() -> bool:
    try:
        import weasyprint as _weasyprint_check
        return bool(_weasyprint_check)
    except ImportError:
        return False


class TestPDFFormatterStreaming:
    """PDFFormatter.begin / append_page / finalize / cleanup."""

    tmpdir = ""
    outfile = ""
    formatter = PDFFormatter(config={})

    def setup_method(self) -> None:
        self.tmpdir = tempfile.mkdtemp()
        self.outfile = os.path.join(self.tmpdir, 'report.pdf')
        self.formatter = PDFFormatter(config={})

    def teardown_method(self) -> None:
        self.formatter.cleanup()
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    @pytest.mark.skipif(not _has_weasyprint(), reason='weasyprint not installed')
    def test_produces_pdf_file(self) -> None:
        """Output starts with %PDF bytes."""
        summary: dict[str, Any] = {'total_pages': 1}
        page_data: dict[str, Any] = {
            'page': _FakePage(url='http://a.com', title='A'),
            'test_result': _FakeTestResult(
                violations=[_FakeIssue(id='ErrNoAlt')],
            ),
        }
        self.formatter.begin(self.outfile, summary)
        self.formatter.append_page(self.outfile, page_data)
        self.formatter.finalize(self.outfile, summary)

        assert os.path.exists(self.outfile)
        with open(self.outfile, 'rb') as f:
            header = f.read(4)
        assert header == b'%PDF'
